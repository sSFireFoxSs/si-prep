"""항목 6-2: 설정 검사 결과 — 프로필 전체 읽기, [검토] 주석 세기, 아직 만들지 않은 방식 경고, 시트별 표.
엑셀은 열지 않는다.
"""

import shutil

import pytest
import yaml

from si_prep.config import ConfigError, check_project, format_check

UNSUPPORTED = "방식은 아직 지원하지 않습니다. 실행 때 실패합니다."


def _copy_shop(tmp_path, generated_dir):
    """shop 설정을 tmp_path로 복사(input_dir은 생성 폴더 절대 경로). project.yaml 경로를 돌려준다"""
    shutil.copytree(SHOP, tmp_path / "shop")
    project = tmp_path / "shop" / "project.yaml"
    data = yaml.safe_load(project.read_text(encoding="utf-8"))
    data["input_dir"] = str(generated_dir)
    project.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return project


def _edit_profile(project, name, mutate):
    path = project.parent / "profiles" / f"{name}.yaml"
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    mutate(data)
    path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return path


def _append(path, text):
    path.write_text(path.read_text(encoding="utf-8") + text, encoding="utf-8")


# ---------- 프로필 전체 읽기 ----------

def test_shop_check_reads_all_profiles(tmp_path, generated_dir):
    result = check_project(_copy_shop(tmp_path, generated_dir))
    assert sorted(result.profiles) == ["bo", "fo", "policy", "req", "test", "variation"]
    assert result.review_marks == {}


def test_errors_from_all_profiles_reported_at_once(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, "fo", lambda d: d.update(heder_rows=[4]))
    _edit_profile(project, "bo", lambda d: d.update(data_start_row=4))
    with pytest.raises(ConfigError) as e:
        check_project(project)
    text = "\n".join(e.value.issues)
    assert "fo.yaml" in text and "heder_rows" in text and "bo.yaml" in text and "data_start_row" in text


# ---------- 아직 만들지 않은 방식: 오류가 아니라 경고 ----------

def test_shop_variation_indent_override_warns(tmp_path, generated_dir):
    # variation의 들여쓰기 시트는 일부러 indent를 쓴다(미구현 → 실행 때 unsupported).
    # 미구현 여부를 등록표에 물으므로, 아직 만들지 않은 split_cell·numbering·
    # continuation 시트도 사실대로 경고한다
    # 아직 만들지 않은 연결 key도 등록표 기준으로 사실대로 경고
    # 2026-10-13: id_ref 구현으로 test_req 경고 없음
    warnings = check_project(_copy_shop(tmp_path, generated_dir)).warnings
    assert warnings == [
        f"variation.yaml sheet_overrides[1] ['한칸경로'] hierarchy: split_cell {UNSUPPORTED}",
        f"variation.yaml sheet_overrides[2] ['번호'] hierarchy: numbering {UNSUPPORTED}",
        f"variation.yaml sheet_overrides[3] ['이어지는행'] item_boundary: continuation {UNSUPPORTED}",
        f"variation.yaml sheet_overrides[4] ['들여쓰기'] hierarchy: indent {UNSUPPORTED}",
    ]


@pytest.mark.parametrize("name,mutate,expected", [
    ("fo", lambda d: d.update(hierarchy={"method": "indent"}), f"fo.yaml hierarchy: indent {UNSUPPORTED}"),
    ("fo", lambda d: d.update(hierarchy={"method": "parent_id"}), f"fo.yaml hierarchy: parent_id {UNSUPPORTED}"),
    ("req", lambda d: d.update(item_boundary={"method": "merged_block"}),
     f"req.yaml item_boundary: merged_block {UNSUPPORTED}"),
])
def test_unimplemented_method_is_warning_not_error(tmp_path, generated_dir, name, mutate, expected):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, name, mutate)
    assert expected in check_project(project).warnings


def test_unimplemented_link_key_warns(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    data = yaml.safe_load(project.read_text(encoding="utf-8"))
    data["link_rules"].append({"id": "req_fo_similar", "docs": ["req", "fo"], "key": "name_similarity"})
    project.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    assert f"project.yaml link_rules (id=req_fo_similar) key: name_similarity {UNSUPPORTED}" in check_project(project).warnings


# ---------- [검토] 주석 ----------

def test_review_marks_counted_per_file(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _append(project.parent / "profiles" / "fo.yaml", "# [검토] 요약 시트가 맞는지\n# [검토] 분류 열 수\n")
    _append(project, "# [검토] 연결 규칙\n")
    result = check_project(project)
    assert result.review_marks == {"project.yaml": 1, "profiles/fo.yaml": 2}


def test_approved_with_review_marks_warns(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)   # shop 프로필은 approved: true
    _append(project.parent / "profiles" / "fo.yaml", "# [검토] 확인\n")
    warnings = check_project(project).warnings
    assert "profiles/fo.yaml: approved: true인데 [검토] 1곳이 남아 있습니다" in warnings


def test_not_approved_with_review_marks_no_warning(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    path = _edit_profile(project, "bo", lambda d: d.update(approved=False))
    _append(path, "# [검토] 확인\n")
    result = check_project(project)
    assert result.review_marks == {"profiles/bo.yaml": 1}
    assert not any("bo.yaml: approved" in w for w in result.warnings)


# ---------- 시트별 표 ----------

def test_sheet_table_rows(tmp_path, generated_dir):
    rows = check_project(_copy_shop(tmp_path, generated_dir)).sheet_rows
    assert len(rows) == 1 + 6 + 2 + 5 + 2 + 24  # 모든 문서의 프로필 sheets 줄
    by_key = {(r.doc_type, r.sheet): r for r in rows}
    cover = by_key[("fo", "표지")]
    assert (cover.read, cover.reason, cover.overview) == (False, "메타 시트(표지)", False)
    assert by_key[("fo", "요약")].overview is True and by_key[("fo", "회원")].overview is False
    assert [r.doc_type for r in rows][:2] == ["req", "fo"]  # documents 순서


def test_format_check_shows_table_and_counts(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _append(project.parent / "profiles" / "fo.yaml", "# [검토] 확인\n")
    text = format_check(check_project(project))
    assert "| 문서 | 시트 | 읽음/뺌 | 이유 | overview | 항목 수 | 링크 수 |" in text
    assert "| fo | 표지 | 뺌 | 메타 시트(표지) |  | (실행 때) | (실행 때) |" in text
    assert "| fo | 요약 | 읽음 |  | 예 | (실행 때) | (실행 때) |" in text
    assert "[검토] 남은 곳: 1 (profiles/fo.yaml 1)" in text
    assert UNSUPPORTED in text  # 경고 목록


# ======================================================================
# 6-3: 설정↔프로필 교차 검사, 문서 항목(파일) 기준 결과, project.yaml 오류와 프로필 오류 함께 보이기
# ======================================================================

def _edit_project(project, mutate):
    data = yaml.safe_load(project.read_text(encoding="utf-8"))
    mutate(data)
    project.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def _rule(data, rule_id):
    return next(r for r in data["link_rules"] if r["id"] == rule_id)


def _check_issues(project):
    with pytest.raises(ConfigError) as e:
        check_project(project)
    return e.value.issues


CROSS_CASES = [
    ("overview를 썼는데 overview_sheets 없음", "fo", lambda d: d.pop("overview_sheets"), None,
     "link_rules[0] (id=fo_overview_detail).sheets[0]", "overview_sheets"),
    ("id_ref인데 참조 열이 어느 쪽에도 없음", "test", lambda d: d["columns"].pop("ref_ids"), None,
     "link_rules[4] (id=test_req)", "ref_ids 열이 어느 쪽"),
    ("id_ref인데 참조 열이 양쪽에", "req", lambda d: d["columns"].update(ref_ids={"header_contains": "설명", "pattern": r"X-\d"}),
     None, "link_rules[4] (id=test_req)", "양쪽"),
    ("id_ref 대상 쪽에 id 열 없음", "req", lambda d: d["columns"].pop("id"), None,
     "link_rules[4] (id=test_req)", "id 열"),
    ("sheet_to_path인데 root_separators 양쪽", "fo", lambda d: d.update(root_separators=[" - "]), None,
     "link_rules[3] (id=policy_fo)", "양쪽"),
    ("sheet_to_path인데 root_separators 어느 쪽에도 없음", "policy", lambda d: d.pop("root_separators"), None,
     "link_rules[3] (id=policy_fo)", "root_separators"),
    ("check_content인데 description 없음", "fo", lambda d: d["columns"].pop("description"), None,
     "link_rules[0] (id=fo_overview_detail)", "description"),
]


@pytest.mark.parametrize("name,profile,mutate,_,where,text", CROSS_CASES, ids=[c[0] for c in CROSS_CASES])
def test_cross_check_with_profiles(tmp_path, generated_dir, name, profile, mutate, _, where, text):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, profile, mutate)
    issues = _check_issues(project)
    assert any(where in m and text in m for m in issues), issues


def _two_fo_documents(tmp_path, generated_dir):
    """fo 문서 항목이 2개(파일·프로필 각각)인 설정. 입력 폴더는 tmp에 복사(생성 폴더는 건드리지 않음)"""
    project = _copy_shop(tmp_path, generated_dir)
    data_dir = tmp_path / "data"
    shutil.copytree(generated_dir, data_dir)
    shutil.copy(data_dir / "가상몰_기능정의서_v1.2.xlsx", data_dir / "가상몰_기능정의서_2판.xlsx")
    shutil.copy(project.parent / "profiles" / "fo.yaml", project.parent / "profiles" / "fo2.yaml")

    def mutate(d):
        d["input_dir"] = str(data_dir)
        d["documents"].insert(2, {"doc_type": "fo", "profile": "profiles/fo2.yaml", "files": ["가상몰_기능정의서_2판.xlsx"]})
    _edit_project(project, mutate)
    return project


def test_two_documents_same_type_errors_both_shown(tmp_path, generated_dir):
    project = _two_fo_documents(tmp_path, generated_dir)
    _edit_profile(project, "fo", lambda d: d.update(heder_rows=[4]))
    _edit_profile(project, "fo2", lambda d: d.update(data_start_row=3))
    text = "\n".join(_check_issues(project))
    assert "fo.yaml" in text and "heder_rows" in text and "fo2.yaml" in text and "data_start_row" in text


def test_two_documents_same_type_sheet_rows_both_shown(tmp_path, generated_dir):
    project = _two_fo_documents(tmp_path, generated_dir)
    _edit_profile(project, "fo2", lambda d: d["sheets"].append({"name": "2판 전용"}))
    result = check_project(project)
    assert [len(d.profile.sheets) for d in result.documents if d.doc_type == "fo"] == [6, 7]
    assert sum(1 for r in result.sheet_rows if r.doc_type == "fo") == 13
    assert sorted(result.profiles) == ["bo", "fo", "policy", "req", "test", "variation"]  # 문서 종류별 묶음
    assert len(result.profiles["fo"]) == 2
    text = format_check(result)
    # 같은 종류 문서가 여럿이면 문서 칸에 파일 이름을 붙여 구분
    assert "| fo (가상몰_기능정의서_2판.xlsx) | 2판 전용 | 읽음 |" in text


def test_project_errors_and_profile_errors_together(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_project(project, lambda d: d["documents"][1].update(doc_type="fo2"))   # project.yaml 오류
    _edit_profile(project, "bo", lambda d: d.update(heder_rows=[4]))              # 프로필 오류
    text = "\n".join(_check_issues(project))
    assert '"fo2"는 doc_types에 선언되지 않았습니다' in text
    assert "bo.yaml" in text and "heder_rows" in text


# ======================================================================
# 항목 6 검토 반영: 시트 이름 맞추기(*·?만 패턴), check_content가 덮어쓰기·sheets 없는 프로필도 봄,
# 규칙이 아무 시트도 못 고르면 경고, 문서 없는 종류 경고, 같은 종류 문서 일부만 열이 있으면 오류,
# 인코딩·권한 오류도 모아 보이기, 같은 메시지 한 번
# ======================================================================

from si_prep.config import Profile, sheet_profile  # noqa: E402


def _mini_profile(overrides):
    return Profile.model_validate({
        "id": "x", "header_rows": [1], "data_start_row": 2,
        "columns": {"hierarchy": [{"header_contains": "분류"}]},
        "hierarchy": {"method": "columns"}, "item_boundary": {"method": "row"},
        "sheet_overrides": overrides})


def test_sheet_name_brackets_are_literal():
    p = _mini_profile([{"sheets": ["[참고]"], "data_start_row": 3}])
    assert sheet_profile(p, "[참고]").data_start_row == 3
    assert sheet_profile(p, "참").data_start_row == 2


def test_sheet_name_star_and_question_still_patterns():
    p = _mini_profile([{"sheets": ["참고*"], "data_start_row": 3}, {"sheets": ["표?"], "data_start_row": 4}])
    assert sheet_profile(p, "참고 1").data_start_row == 3
    assert sheet_profile(p, "표A").data_start_row == 4
    assert sheet_profile(p, "표AB").data_start_row == 2


def test_check_content_override_description_missing_without_sheets_list(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)

    def mutate(d):
        d.pop("sheets")
        d["sheet_overrides"][1]["columns"].pop("description")   # 요약 시트 덮어쓰기
    _edit_profile(project, "fo", mutate)
    issues = _check_issues(project)
    assert any("fo_overview_detail" in m and "description" in m and "sheet_overrides[1]" in m for m in issues), issues


def test_check_content_detail_override_description_missing(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, "fo", lambda d: d["sheet_overrides"][0]["columns"].pop("description"))   # 고객센터
    issues = _check_issues(project)
    assert any("fo_overview_detail" in m and "sheet_overrides[0]" in m for m in issues), issues


def test_overview_sheets_matching_no_sheet_warns(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, "fo", lambda d: d.update(overview_sheets=["요악"]))
    warnings = check_project(project).warnings
    assert any("fo_overview_detail" in w and "overview" in w and "고르는 시트가 없습니다" in w for w in warnings), warnings


def test_rule_sheet_name_matching_no_sheet_warns(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_project(project, lambda d: _rule(d, "bo_test").update(sheets=[None, ["BO테스트캐이스"]]))
    warnings = check_project(project).warnings
    assert any("bo_test" in w and "BO테스트캐이스" in w for w in warnings), warnings


def test_rule_with_doc_type_without_documents_warns(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)

    def mutate(d):
        d["doc_types"].append({"id": "extra", "name": "문서 없는 종류"})
        d["link_rules"].append({"id": "extra_fo", "docs": ["extra", "fo"], "key": "path"})
    _edit_project(project, mutate)
    warnings = check_project(project).warnings
    assert any("extra_fo" in w and "extra" in w and "문서 항목이 없" in w for w in warnings), warnings


def test_same_type_documents_partly_have_ref_ids_is_error(tmp_path, generated_dir):
    project = _two_fo_documents(tmp_path, generated_dir)
    _edit_profile(project, "fo2", lambda d: d["columns"].update(ref_ids={"header_contains": "설명", "pattern": r"R-\d"}))
    _edit_project(project, lambda d: d["link_rules"].append({"id": "fo_req", "docs": ["fo", "req"], "key": "id_ref"}))
    issues = _check_issues(project)
    assert any("fo_req" in m and "일부" in m and "ref_ids" in m for m in issues), issues


def test_profile_not_utf8_reported_with_other_errors(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    fo = project.parent / "profiles" / "fo.yaml"
    fo.write_bytes(fo.read_text(encoding="utf-8").encode("cp949"))
    _edit_profile(project, "bo", lambda d: d.update(heder_rows=[4]))
    text = "\n".join(_check_issues(project))
    assert "fo.yaml" in text and "UTF-8" in text
    assert "bo.yaml" in text and "heder_rows" in text


def test_project_not_utf8_is_config_error(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    project.write_bytes(project.read_text(encoding="utf-8").encode("cp949"))
    assert any("UTF-8" in m for m in _check_issues(project))


def test_shared_profile_cross_error_reported_once(tmp_path, generated_dir):
    project = _two_fo_documents(tmp_path, generated_dir)
    _edit_project(project, lambda d: d["documents"][2].update(profile="profiles/fo.yaml"))   # 두 fo 문서가 같은 프로필
    _edit_profile(project, "fo", lambda d: d.pop("overview_sheets"))
    issues = _check_issues(project)
    assert sum("overview_sheets" in m for m in issues) == 1, issues


def test_missing_documents_side_only_warns_for_id_ref(tmp_path, generated_dir):
    # 문서가 아직 없는 종류: 경고만, ref_ids 「어느 쪽에도 없음」 오류로 실행을 막지 않는다
    project = _copy_shop(tmp_path, generated_dir)
    _edit_project(project, lambda d: d.update(documents=[x for x in d["documents"] if x["doc_type"] != "test"]))
    warnings = check_project(project).warnings
    assert any("test_req" in w and "문서 항목이 없" in w for w in warnings), warnings


def test_check_content_base_not_needed_when_named_sheets_all_overridden(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)

    def mutate(d):
        d["link_rules"] = [r for r in d["link_rules"] if r["id"] != "fo_overview_detail"]
        d["link_rules"].append({"id": "fo_cs_summary", "docs": ["fo", "fo"], "key": "path",
                                "sheets": [["고객센터"], ["요약"]], "check_content": True})
    _edit_project(project, mutate)
    _edit_profile(project, "fo", lambda d: (d.pop("sheets"), d["columns"].pop("description")))
    check_project(project)   # 두 칸 모두 덮어쓰기(description 있음)만 쓰이므로 오류 없음


def test_sheet_table_overview_uses_same_rule_as_cross_check(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, "fo", lambda d: d["sheets"][5].update(name="요약*"))
    rows = check_project(project).sheet_rows
    assert next(r for r in rows if r.sheet == "요약*").overview is True


def test_cross_warnings_start_with_file_name(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_project(project, lambda d: _rule(d, "bo_test").update(sheets=[None, ["BO테스트캐이스"]]))
    warnings = [w for w in check_project(project).warnings if "bo_test" in w]
    assert warnings and all(w.startswith("project.yaml ") for w in warnings)


def test_missing_documents_side_still_checks_other_side(tmp_path, generated_dir):
    # 한쪽 문서가 없어도 있는 쪽의 overview·check_content 검사는 한다
    project = _copy_shop(tmp_path, generated_dir)

    def mutate(d):
        d["doc_types"].append({"id": "extra", "name": "문서 없는 종류"})
        d["link_rules"].append({"id": "fo_extra", "docs": ["fo", "extra"], "key": "path",
                                "sheets": ["overview", None], "check_content": True})
    _edit_project(project, mutate)
    _edit_profile(project, "fo", lambda d: d.pop("overview_sheets"))
    assert any("fo_extra" in m and "overview_sheets" in m for m in _check_issues(project))


# ======================================================================
# 6c: 엑셀 시트 대조 (read_excel=True 또는 --excel 일 때만 엑셀을 연다, 시트 이름만)
# ======================================================================

import openpyxl  # noqa: E402

from si_prep import reader as _reader  # noqa: E402
from si_prep.config import main as config_main  # noqa: E402
from conftest import SHOP_DIR as SHOP

FO_FILE = "가상몰_기능정의서_v1.2.xlsx"
NEW_SHEET = "새 시트, 확인 필요"


def _excel_shop(tmp_path, generated_dir):
    """--excel 검사용 shop: 변형 문서는 뺀다(일부러 넣은 머리글 실패 시트가 9a 설정 검사에서 오류가 되므로)"""
    project = _copy_shop(tmp_path, generated_dir)
    _edit_project(project, lambda d: d.update(documents=[x for x in d["documents"] if x["doc_type"] != "variation"]))
    return project


def _shop_with_data_copy(tmp_path, generated_dir):
    """입력 폴더를 tmp로 복사한 shop 설정(엑셀을 고쳐도 생성 폴더는 그대로), 변형 문서 뺌"""
    project = _excel_shop(tmp_path, generated_dir)
    data_dir = tmp_path / "data"
    shutil.copytree(generated_dir, data_dir)
    _edit_project(project, lambda d: d.update(input_dir=str(data_dir)))
    return project, data_dir


def _add_sheet(path, name):
    wb = openpyxl.load_workbook(path)
    wb.create_sheet(name)["A1"] = "값"
    wb.save(path)


def test_without_read_excel_no_excel_opened(tmp_path, generated_dir, monkeypatch):
    def boom(*a, **k):
        raise AssertionError("엑셀을 열면 안 됨")
    for name in ("sheet_names", "sheet_list", "open_workbook"):
        monkeypatch.setattr(_reader, name, boom)
    check_project(_copy_shop(tmp_path, generated_dir))


def test_shop_excel_check_without_variation_no_warnings(tmp_path, generated_dir):
    # 변형 문서(일부러 넣은 실패 시트·미지원 파일)를 뺀 shop은 경고 없음
    result = check_project(_excel_shop(tmp_path, generated_dir), read_excel=True)
    # 2026-10-10: 남는 경고는 아직 만들지 않은 연결 key뿐 → 2026-10-13 id_ref 구현으로 0개
    assert result.warnings == []
    assert not any(r.new for r in result.sheet_rows)


def test_new_sheet_in_excel_warns_once(tmp_path, generated_dir):
    project, data_dir = _shop_with_data_copy(tmp_path, generated_dir)
    _add_sheet(data_dir / FO_FILE, "이벤트")
    base = check_project(_excel_shop(tmp_path / "b", generated_dir), read_excel=True).warnings
    result = check_project(project, read_excel=True)
    new = [w for w in result.warnings if w not in base]
    assert len(new) == 1 and NEW_SHEET in new[0] and "이벤트" in new[0] and FO_FILE in new[0], new
    row = next(r for r in result.sheet_rows if r.sheet == "이벤트")
    assert (row.doc_type, row.read, row.new) == ("fo", True, True)
    assert f"| fo | 이벤트 | 읽음 · {NEW_SHEET} |" in format_check(result)


def test_new_sheet_matching_pattern_is_not_new(tmp_path, generated_dir):
    project, data_dir = _shop_with_data_copy(tmp_path, generated_dir)
    _add_sheet(data_dir / FO_FILE, "참고 1")
    _edit_profile(project, "fo", lambda d: d["sheets"].append({"name": "참고*", "read": False, "reason": "참고 시트"}))
    result = check_project(project, read_excel=True)
    assert not any(NEW_SHEET in w for w in result.warnings)


def test_profile_sheet_missing_in_excel_warns(tmp_path, generated_dir):
    project = _excel_shop(tmp_path, generated_dir)
    _edit_profile(project, "fo", lambda d: d["sheets"].append({"name": "없는시트"}))
    result = check_project(project, read_excel=True)
    hits = [w for w in result.warnings if "없는시트" in w]
    assert len(hits) == 1 and "엑셀에 없습니다" in hits[0] and FO_FILE in hits[0], hits


def test_profile_without_sheets_list_has_no_new_sheets(tmp_path, generated_dir):
    project, data_dir = _shop_with_data_copy(tmp_path, generated_dir)
    _add_sheet(data_dir / FO_FILE, "이벤트")
    _edit_profile(project, "fo", lambda d: d.pop("sheets"))    # 목록이 없으면 모든 시트를 읽음 → 「새 시트」 아님
    # 모든 시트를 읽으므로 머리글 없는 표지·개정이력·새 시트는 머리글 오류
    issues = _check_issues_excel(project)
    assert not any(NEW_SHEET in m for m in issues)
    assert any('"표지"' in m for m in issues) and any('"이벤트"' in m for m in issues)


def test_cli_excel_flag(tmp_path, generated_dir, capsys):
    project, data_dir = _shop_with_data_copy(tmp_path, generated_dir)
    _add_sheet(data_dir / FO_FILE, "이벤트")
    assert config_main([str(project)]) == 0
    assert NEW_SHEET not in capsys.readouterr().out          # 기본은 엑셀을 열지 않음
    assert config_main([str(project), "--excel"]) == 0
    assert NEW_SHEET in capsys.readouterr().out


def test_new_hidden_sheet_says_skipped_not_read(tmp_path, generated_dir):
    project, data_dir = _shop_with_data_copy(tmp_path, generated_dir)
    wb = openpyxl.load_workbook(data_dir / FO_FILE)
    ws = wb.create_sheet("임시")
    ws["A1"] = "값"
    ws.sheet_state = "hidden"
    wb.save(data_dir / FO_FILE)
    result = check_project(project, read_excel=True)
    [w] = [w for w in result.warnings if "임시" in w]
    assert NEW_SHEET in w and "숨김 시트" in w and "건너뜀" in w and "읽습니다" not in w
    row = next(r for r in result.sheet_rows if r.sheet == "임시")
    assert (row.read, row.new, row.reason) == (False, True, "숨김 시트(실행 때 건너뜀)")


def test_new_visible_sheet_message_mentions_profile_settings(tmp_path, generated_dir):
    project, data_dir = _shop_with_data_copy(tmp_path, generated_dir)
    _add_sheet(data_dir / FO_FILE, "이벤트")
    [w] = [w for w in check_project(project, read_excel=True).warnings if "이벤트" in w]
    assert "프로필 설정(시트별 덮어쓰기 포함)대로 읽습니다" in w


def test_pattern_line_without_match_says_not_yet(tmp_path, generated_dir):
    project = _excel_shop(tmp_path, generated_dir)
    _edit_profile(project, "fo", lambda d: d["sheets"].append({"name": "참고*", "read": False, "reason": "참고 시트"}))
    [w] = [w for w in check_project(project, read_excel=True).warnings if "참고*" in w]
    assert "패턴" in w and "아직 맞는 시트가 없습니다" in w


# ======================================================================
# 9a: --excel 설정 검사의 머리글 대조 (읽음 시트가 머리글을 못 찾으면 오류, 선택자별 맞은 열 표)
# ======================================================================

def test_full_shop_excel_reports_intended_broken_sheets_as_errors(tmp_path, generated_dir):
    issues = _check_issues_excel(_copy_shop(tmp_path, generated_dir))
    for sheet in ("역할겹침", "분류열부족", "빈시트"):
        assert any(sheet in m and "가상몰_변형.xlsx" in m for m in issues), (sheet, issues)
    assert not any("들여쓰기" in m for m in issues)          # 미구현 방식 시트는 실행 때도 건너뛰므로 검사 안 함


def _check_issues_excel(project):
    with pytest.raises(ConfigError) as e:
        check_project(project, read_excel=True)
    return e.value.issues


def test_header_report_rows(tmp_path, generated_dir):
    text = format_check(check_project(_excel_shop(tmp_path, generated_dir), read_excel=True))
    assert "| 문서 | 시트 | 머리글 열 수 | 역할 | 맞은 열 수 | 열: 합친 머리글 |" in text
    assert "| fo | 회원 | 5 | title | 1 | E: 기능명 |" in text                       # 머리글 = 선택자 글자 그대로
    # 2026-10-09 작성자: header_contains로 맞은 열은 합친 머리글 전체를 보인다
    assert "| fo | 회원 | 5 | hierarchy | 3 | B: 분류 / 1단계, C: 분류 / 2단계, D: 분류 / 3단계 |" in text


def test_header_report_hides_text_of_col_selected_columns(tmp_path, generated_dir):
    # 열 문자로 지정한 열은 머리글 행 설정이 틀리면 데이터 칸일 수 있어 글자를 보이지 않는다
    text = format_check(check_project(_excel_shop(tmp_path, generated_dir), read_excel=True))
    assert "| policy | 주문 - 배송 |" in text
    # 2026-10-09 작성자: 2칸 이상 가로 병합된 머리글 칸(묶음 제목)은 보인다(B1:C1 「세부 구분」)
    assert "hierarchy | 2 | B: 세부 구분, C: 세부 구분 |" in text


def test_id_ref_target_needs_id_column_in_excel(tmp_path, generated_dir):
    project = _excel_shop(tmp_path, generated_dir)
    _edit_profile(project, "req", lambda d: d["columns"].update(id={"header_contains": "없는 머리글"}))
    issues = _check_issues_excel(project)
    assert any("가상몰_요구사항정의서.xlsx" in m and "id" in m for m in issues), issues


def test_header_report_sheet_without_roles_still_listed(tmp_path, generated_dir):
    # 역할을 지정하지 않은 시트(정책 목록형 등)도 머리글을 찾았는지·머리글 열 수는 보인다(머리글 글자는 안 보임)
    project = _excel_shop(tmp_path, generated_dir)

    def mutate(d):
        d["columns"] = {}
        for o in d["sheet_overrides"]:
            if o["sheets"] == ["결제 수단 목록"]:
                o.pop("columns")
    _edit_profile(project, "policy", mutate)
    text = format_check(check_project(project, read_excel=True))
    assert any(line.startswith("| policy | 결제 수단 목록 |") and "| (역할 지정 없음) | 0 |  |" in line
               for line in text.splitlines()), text


def test_header_report_pattern_shows_text_only_on_full_match(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_project(project, lambda d: d.update(documents=[x for x in d["documents"] if x["doc_type"] == "variation"],
                                             link_rules=[]))
    _edit_profile(project, "variation", lambda d: d.update(sheets=[e for e in d["sheets"] if e["name"] in ("분류패턴_2단",)]))
    text = format_check(check_project(project, read_excel=True))
    assert "| variation | 분류패턴_2단 | 4 | hierarchy | 2 | B: 분류1, C: 분류2 |" in text


def test_optional_selector_not_found_warns_in_excel_check(tmp_path, generated_dir):
    project = _excel_shop(tmp_path, generated_dir)
    _edit_profile(project, "bo", lambda d: d["columns"].update(title={"header_contains": "없는 머리글"}))
    warnings = check_project(project, read_excel=True).warnings
    assert any("가상몰_관리자기능정의서.xlsx" in w and "title" in w and "못 찾" in w for w in warnings), warnings


def test_id_required_only_on_rule_selected_sheets(tmp_path, generated_dir):
    # id_ref 대상 쪽이라도 규칙 sheets 칸이 고르지 않는 시트는 id 열을 요구하지 않는다
    project = _excel_shop(tmp_path, generated_dir)
    _edit_project(project, lambda d: _rule(d, "test_req").update(sheets=[None, ["다른시트"]]))
    _edit_profile(project, "req", lambda d: d["columns"].update(id={"header_contains": "없는 머리글"}))
    check_project(project, read_excel=True)     # 요구사항 시트는 규칙 밖 → 오류 없음


def test_excel_errors_also_show_excel_warnings(tmp_path, generated_dir):
    project, data_dir = _shop_with_data_copy(tmp_path, generated_dir)
    _add_sheet(data_dir / FO_FILE, "이벤트")
    _edit_profile(project, "req", lambda d: d["columns"].update(id={"header_contains": "없는 머리글"}))
    issues = _check_issues_excel(project)
    assert any(m.startswith("(경고) ") and NEW_SHEET in m for m in issues), issues


def test_header_report_escapes_pipe():
    from si_prep.config import _cell_text
    assert _cell_text("가|나") == r"가\|나"


@pytest.mark.parametrize("sel,parts,shown", [
    ({"header_contains": "정의"}, [("기능정의", True), ("정의", False)], "F: 기능정의 / 정의"),
    ({"col": "F"}, [("세부 구분", True)], "F: 세부 구분"),
    ({"col": "F"}, [("수행 결과", True), ("테스터", False)], "F: 수행 결과 (+1줄 글자 가림: 병합 안 된 머리글 칸)"),
    ({"col": "F"}, [("값", False)], "F (열 지정)"),
    ({"header_pattern": "^분류$", "all": True}, [("묶음", True), ("분류", False)], "F: 묶음 (+1줄 글자 가림: 병합 안 된 머리글 칸)"),
    ({"header_pattern": "분류", "all": True}, [("분류", False)], "F: 분류"),
])
def test_header_shown_rules(sel, parts, shown):
    from si_prep.config import Selector, _shown
    header = " / ".join(t for t, _ in parts)
    assert _shown("F", header, Selector.model_validate(sel), parts) == shown


# ---------- 2026-10-10: 버전을 못 찾으면 --excel 설정 검사 경고 ----------

@pytest.mark.parametrize("change,where", [
    ({"filename_pattern": r"_ver(\d+)"}, "파일 이름"),
    ({"label_pattern": "없는 라벨"}, "표지"),
])
def test_version_not_found_warns_in_excel_check(tmp_path, generated_dir, change, where):
    project = _excel_shop(tmp_path, generated_dir)
    _edit_profile(project, "fo", lambda d: d["version"].update(change))
    warnings = check_project(project, read_excel=True).warnings
    assert any(FO_FILE in w and "버전을 못 찾" in w and where in w for w in warnings), warnings


def test_version_found_no_warning(tmp_path, generated_dir):
    warnings = check_project(_excel_shop(tmp_path, generated_dir), read_excel=True).warnings
    assert not any("버전을 못 찾" in w for w in warnings)


# ---------- 2026-10-10 (15 검토 M3): sheet_to_path 찾는 쪽 시트는 sheet_as_root: true ----------

def test_sheet_to_path_seeker_needs_sheet_as_root(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, "policy", lambda d: d.update(sheet_as_root=False))
    issues = _check_issues(project)
    assert any("policy_fo" in m and "sheet_as_root" in m for m in issues), issues


def test_sheet_to_path_seeker_override_off_outside_rule_slot_is_ok(tmp_path, generated_dir):
    # 규칙 칸(["* - *"]) 밖 시트(결제 수단 목록)는 sheet_as_root가 꺼져도 상관없음
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, "policy", lambda d: d["sheet_overrides"].append(
        {"sheets": ["결제 수단 목록"], "sheet_as_root": False}))
    check_project(project)


# ---------- 2026-10-12 리뷰 R19: 프로필에 sheets 목록이 없을 때 sheet_as_root 판정(기본값·덮어쓰기) ----------

def test_sheet_to_path_seeker_without_sheets_list_default_off(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, "policy", lambda d: (d.pop("sheets"), d.update(sheet_as_root=False)))
    issues = _check_issues(project)
    assert any("policy_fo" in m and "기본값" in m and "sheet_as_root" in m for m in issues), issues


def test_sheet_to_path_seeker_without_sheets_list_override_off_in_slot(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, "policy", lambda d: (d.pop("sheets"), d["sheet_overrides"].append(
        {"sheets": ["주문 - *"], "sheet_as_root": False})))
    issues = _check_issues(project)
    assert any("policy_fo" in m and "sheet_overrides[" in m and "sheet_as_root" in m for m in issues), issues


def test_sheet_to_path_seeker_without_sheets_list_override_off_outside_slot_ok(tmp_path, generated_dir):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_profile(project, "policy", lambda d: (d.pop("sheets"), d["sheet_overrides"].append(
        {"sheets": ["결제 수단 목록"], "sheet_as_root": False})))
    check_project(project)


# ---------- 아직 구현 안 된 설정 키는 「효과 없음」 경고 ----------

UNIMPLEMENTED_KEY = "설정 키는 아직 구현되지 않아 효과가 없습니다"


@pytest.mark.parametrize("mutate,expected", [
    (lambda d: d.update(table_gap_rows=3), f"fo.yaml table_gap_rows: {UNIMPLEMENTED_KEY}"),
    (lambda d: d.update(table_gap_rows=2), f"fo.yaml table_gap_rows: {UNIMPLEMENTED_KEY}"),   # 기본값과 같아도 적었으면 경고
    (lambda d: d["sheet_overrides"].append({"sheets": ["회원"], "table_gap_rows": 1}),
     f"fo.yaml sheet_overrides[{{n}}] ['회원'] table_gap_rows: {UNIMPLEMENTED_KEY}"),
])
def test_unimplemented_key_warns(tmp_path, generated_dir, mutate, expected):
    project = _copy_shop(tmp_path, generated_dir)
    path = _edit_profile(project, "fo", mutate)
    n = len(yaml.safe_load(path.read_text(encoding="utf-8"))["sheet_overrides"]) - 1
    warnings = check_project(project).warnings
    [w] = [w for w in warnings if w.startswith(expected.format(n=n))]
    assert "시트 건너뜀" in w                                # 무엇이 없는지(설계상 동작)도 같이


def test_unimplemented_key_absent_no_warning(tmp_path, generated_dir):
    warnings = check_project(_copy_shop(tmp_path, generated_dir)).warnings
    assert not any(UNIMPLEMENTED_KEY in w for w in warnings)


# ---------- 머리글 대조에 숨김 열·값 없는 열 표시 ----------

def _marks_table(tmp_path):
    import openpyxl
    from si_prep.config import Profile, sheet_profile
    from si_prep.issues import Issues
    from si_prep.layout import build_table
    from si_prep.reader import open_workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r, row in enumerate([["분류", "기능명", "설명", "비고"], ["가", "x", None, "m"], ["나", "y", None, None]], start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
    ws.column_dimensions["C"].hidden = True       # 숨김 + 값 없음
    ws.column_dimensions["D"].hidden = True       # 숨김, 값 있음
    path = tmp_path / "m.xlsx"
    wb.save(path)
    p = Profile.model_validate({"id": "x", "header_rows": [1], "data_start_row": 2,
                                "columns": {"hierarchy": [{"header_contains": "분류"}], "title": {"header_contains": "기능명"},
                                            "description": {"header_contains": "설명"}},
                                "hierarchy": {"method": "columns"}, "item_boundary": {"method": "row"}})
    return build_table(open_workbook(path, Issues()).sheet("S"), sheet_profile(p, "S"), "m.xlsx", Issues())


def test_header_column_marks_hidden_and_empty(tmp_path):
    from si_prep.config import _column_marks
    table = _marks_table(tmp_path)
    assert _column_marks(table, "A") == ""
    assert _column_marks(table, "C") == " [숨김 열, 값 없음]"
    assert _column_marks(table, "D") == " [숨김 열]"


def test_header_report_explains_marks(tmp_path, generated_dir):
    text = format_check(check_project(_excel_shop(tmp_path, generated_dir), read_excel=True))
    assert "[숨김 열] = 엑셀에서 숨긴 열" in text and "[값 없음] = 데이터 행에 원문 값이 없는 열" in text


# ---------- 9b: tables 시트는 표마다 머리글 대조, 표 하나 건너뜀은 경고, 표가 없으면 오류 ----------

def _tables_only(tmp_path, generated_dir, mutate=None):
    project = _copy_shop(tmp_path, generated_dir)
    _edit_project(project, lambda d: d.update(documents=[x for x in d["documents"] if x["doc_type"] == "variation"],
                                             link_rules=[]))

    def keep(d):
        d["sheets"] = [e for e in d["sheets"] if e["name"] == "표여러개"]
        if mutate:
            mutate(d)
    _edit_profile(project, "variation", keep)
    return project


def test_excel_check_tables_sheet_per_table(tmp_path, generated_dir):
    result = check_project(_tables_only(tmp_path, generated_dir), read_excel=True)
    text = format_check(result)
    rows = [l for l in text.splitlines() if l.startswith("| variation | 표여러개 [표 제목 ")]
    assert len([l for l in rows if "| hierarchy |" in l]) == 4                  # 표 4개 대조(표3은 건너뜀)
    assert any("표(제목 11행)" in w and "이 표를 건너뜁니다" in w for w in result.warnings)


def test_excel_check_tables_sheet_without_tables_is_error(tmp_path, generated_dir):
    def bad(d):
        o = next(o for o in d["sheet_overrides"] if o["sheets"] == ["표여러개"])
        o["tables"]["header_row"]["pattern"] = "^없는머리글$"
    issues = _check_issues_excel(_tables_only(tmp_path, generated_dir, bad))
    assert any("표여러개" in m and "tables" in m for m in issues), issues


# ---------- 표마다 대조 줄에 「머리글 n줄」 ----------

def test_excel_check_tables_label_shows_header_lines(tmp_path, generated_dir):
    text = format_check(check_project(_tables_only(tmp_path, generated_dir), read_excel=True))
    rows = [l for l in text.splitlines() if l.startswith("| variation | 표여러개 [표 제목 ") and "| hierarchy |" in l]
    assert rows and all(" · 머리글 1줄] |" in l for l in rows)
    assert "| variation | 표여러개 [표 제목 1행 · 머리글 1줄] |" in text
