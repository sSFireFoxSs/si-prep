"""실제 용어 겹침 검사 (합성 fixture가 쓴다).

실제 용어는 진단표(docs/intake/*.md, 커밋 안 됨)에서 실행할 때 읽는다. 코드에 복사하지 않는다.
실패해도 실제 용어를 출력하지 않는다(개수만).
"""

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

INTAKE_DIR = Path(__file__).resolve().parents[1] / "docs" / "intake"

# 허용 단어 6개: 문서 양식에 공통으로 쓰는 일반 단어이고, docs/DESIGN.md 5·8장의
# 합성 예시에 이미 공개돼 있다. 진단표의 실제 용어에도 같은 글자로 있어서 예외로 둔다.
#   - "표지": 표지 시트 이름 (5장 fo 프로필 version.cover_sheet 예시)
#   - "설명": 설명 열 머리글 (5장 fo·req 프로필 columns.description 예시)
#   - "케이스": 테스트 케이스 문서의 일반 단어 (5장 test 프로필 "케이스 ID" 예시)
#   - "내용": 일반 단어 (8.1 시나리오 시트 이름 "내용병합")
#   - "NO.": 맨 앞 번호 열 머리글 (8.1 「고객센터」 시트, 실제 양식 구조를 따르기 위해)
#   - "비고": 대부분 빈 비고 열 머리글 (8.1 「고객센터」 시트, 같은 이유)
# 메뉴·기능·회사·사람 이름처럼 보이는 단어는 넣지 않는다. 실제로 합성 데이터에 쓰는 단어만 넣는다.
ALLOWED_GENERIC_WORDS = {"표지", "설명", "케이스", "내용", "NO.", "비고"}


def _section(text, heading):
    """'## heading'부터 다음 '## '까지."""
    m = re.search(rf"^## {re.escape(heading)}.*?$(.*?)(?=^## |\Z)", text, re.M | re.S)
    return m.group(1) if m else ""


def has_intake():
    return any(p.name != "README.md" for p in INTAKE_DIR.glob("*.md"))


def load_real_terms():
    """진단표의 「실제 용어」 절의 용어만. 「열 이름과 역할」 표의 열 이름은 흔한 양식 단어라 오탐이 나서 읽지 않는다
   """
    terms = set()
    for md in sorted(INTAKE_DIR.glob("*.md")):
        if md.name == "README.md":
            continue
        text = md.read_text(encoding="utf-8")
        for line in _section(text, "실제 용어").splitlines():
            m = re.match(r"- [^:]+: (.*)", line)
            if m:
                terms |= set(_terms(m.group(1)))
    return {t for t in terms if t}


def _terms(text):
    """「용어, 용어, …」 → 용어 목록 (intake README 「테스트가 읽는 절」).
    쉼표+공백으로 나누되 백틱 안은 나누지 않는다. 백틱으로 감싼 용어는 그 안 글자 그대로.
    그 밖에는 공백 뒤 괄호만 설명으로 뺀다(공백 없이 붙은 괄호는 이름의 일부, 예: 시트 이름 「목록(관리자)」)"""
    for token in re.findall(r"(?:`[^`]*`|[^,`]|,(?! ))+", text):
        quoted = re.match(r"\s*`([^`]*)`", token)
        yield quoted.group(1) if quoted else re.sub(r"\s+\(.*\)$", "", token).strip()


def xlsx_strings(paths):
    """엑셀 파일 이름·시트 이름·글자 칸 값. 옛 형식(zip 아님) 파일은 이름만"""
    strings = set()
    for path in sorted(paths):
        strings.add(Path(path).stem)
        if not zipfile.is_zipfile(path):
            continue
        wb = openpyxl.load_workbook(path)
        for name in wb.sheetnames:
            strings.add(name)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                strings |= {str(c.value).strip() for c in row if isinstance(c.value, str)}
    return strings


def nested_strings(node, strings=None):
    """dict·list 안의 키와 글자 값 전부 (yaml·json 검사용)"""
    strings = set() if strings is None else strings
    if isinstance(node, dict):
        for k, v in node.items():
            strings.add(str(k))
            nested_strings(v, strings)
    elif isinstance(node, list):
        for v in node:
            nested_strings(v, strings)
    elif isinstance(node, str):
        strings.add(node.strip())
    return strings


def check_no_real_terms(strings, what="합성 데이터"):
    """진단표가 없으면 skip, 실제 용어가 섞였으면 개수만 보여 주고 실패"""
    if not has_intake():
        pytest.skip("진단표(docs/intake/*.md)가 없는 환경: 실제 용어 검사를 할 수 없음")
    real = load_real_terms()
    assert real, "진단표는 있는데 실제 용어를 하나도 읽지 못함 (진단표 양식 확인)"
    checked = real - ALLOWED_GENERIC_WORDS
    # 허용 단어와 똑같은 합성 글자는 검사하지 않는다. 그 안에 다른 실제 용어가
    # 부분으로 들어 있어도 허용 단어 자체를 쓴 것이므로
    candidates = set(strings) - ALLOWED_GENERIC_WORDS
    # 글자 전체가 같거나, 실제 용어가 그 안에 들어 있으면 실패
    hits = {t for t in checked for s in candidates if t == s or (len(t) >= 2 and t in s)}
    # 실패 메시지에 실제 용어를 그대로 쓰지 않는다(개수만). assert를 쓰면 pytest가 목록을 출력하므로 fail로
    if hits:
        pytest.fail(f"{what}에 실제 용어 {len(hits)}개가 섞임", pytrace=False)
