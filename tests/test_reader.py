"""reader — openpyxl을 쓰는 유일한 모듈. 칸 값·출처·숨김·수식·파일 형식.

생성된 합성 파일(가상몰_변형.xlsx 등)과, 테스트 안에서 바로 만드는 작은 엑셀을 쓴다. 실제 문서는 쓰지 않는다.
"""

import ast
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
import pytest
from openpyxl.comments import Comment

from si_prep import reader
from si_prep.issues import Issues
from si_prep.reader import open_workbook

SRC = Path(__file__).resolve().parents[1] / "src" / "si_prep"


def _save(tmp_path, build, name="t.xlsx"):
    wb = openpyxl.Workbook()
    build(wb.active)
    path = tmp_path / name
    wb.save(path)
    return path


def _open(path):
    issues = Issues()
    return open_workbook(path, issues), issues


# ---------- 병합·빈칸 ----------

def _merged_sheet(ws):
    ws.title = "S"
    ws["B6"] = "로그인"
    ws.merge_cells("B6:B7")
    ws["D6"] = "가"
    ws["D7"] = "나"


def test_merged_blank_is_merge_restored(tmp_path):
    wb, issues = _open(_save(tmp_path, _merged_sheet))
    s = wb.sheet("S")
    b7 = s.cell(7, "B")
    assert (b7.value, b7.origin, b7.row, b7.col) == ("로그인", "merge_restored", 7, "B")
    assert b7.merge_start == (6, "B")
    b6 = s.cell(6, "B")
    assert (b6.value, b6.origin, b6.merge_start) == ("로그인", "cell", (6, "B"))
    assert s.cell(7, "D").origin == "cell" and s.cell(7, "D").merge_start is None
    assert len(issues) == 0


def test_plain_blank_is_empty(tmp_path):
    wb, _ = _open(_save(tmp_path, _merged_sheet))
    c = wb.sheet("S").cell(7, "C")
    assert (c.value, c.origin, c.merge_start) == (None, "empty", None)


def test_merge_with_empty_start_is_empty(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["A1"] = "x"
        ws.merge_cells("B2:B3")              # 시작 칸이 빈 병합
    wb, _ = _open(_save(tmp_path, build))
    c = wb.sheet("S").cell(3, "B")
    assert (c.value, c.origin, c.merge_start) == (None, "empty", (2, "B"))


def test_merge_index_built_once_per_sheet(tmp_path, monkeypatch):
    calls = []
    original = reader._merge_index
    monkeypatch.setattr(reader, "_merge_index", lambda ws: calls.append(ws.title) or original(ws))
    wb, _ = _open(_save(tmp_path, _merged_sheet))
    s = wb.sheet("S")
    for r in range(1, s.max_row + 1):
        for col in s.columns:
            s.cell(r, col)
    assert calls == ["S"]


def test_values_kept_as_is(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["A1"] = "  앞뒤 공백 "
        ws["B1"] = 7
        ws["C1"] = 1.5
        ws["D1"] = "줄\n바꿈"
    s = _open(_save(tmp_path, build))[0].sheet("S")
    assert [s.cell(1, c).value for c in "ABCD"] == ["  앞뒤 공백 ", 7, 1.5, "줄\n바꿈"]


# ---------- 날짜 ----------

def test_dates_become_iso_strings(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["A1"] = datetime(2026, 10, 4)
        ws["B1"] = datetime(2026, 10, 4, 9, 30)
        ws["C1"] = date(2026, 10, 4)
        ws["D1"] = time(9, 30)
    s = _open(_save(tmp_path, build))[0].sheet("S")
    assert [s.cell(1, c).value for c in "ABCD"] == ["2026-10-04", "2026-10-04T09:30:00", "2026-10-04", "09:30:00"]


# ---------- 숨김 ----------

def test_hidden_sheet_rows_cols(generated_dir):
    wb, _ = _open(generated_dir / "가상몰_변형.xlsx")
    assert wb.sheet("숨김").hidden is True
    assert wb.sheet("빈칸채움").hidden is False
    hrf = wb.sheet("숨김행수식")
    assert hrf.hidden_rows == {3: "hidden", 4: "outline"}   # 그냥 숨김 / 그룹 접기
    assert hrf.cell(3, "E").hidden is True and hrf.cell(2, "E").hidden is False


def test_hidden_columns_including_ranges(tmp_path):
    def build(ws):
        ws.title = "S"
        for col in "ABCDEF":
            ws[f"{col}1"] = col
        ws.column_dimensions["B"].hidden = True
        ws.column_dimensions.group("D", "E", hidden=True)   # 범위 하나로 저장되는 숨김 열
    s = _open(_save(tmp_path, build))[0].sheet("S")
    assert s.hidden_cols == ("B", "D", "E")      # 열 순서로 정렬(결정성)
    assert [s.cell(1, c).hidden for c in "ABCDEF"] == [False, True, False, True, True, False]


def test_very_hidden_sheet_is_hidden(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["A1"] = 1
        ws.sheet_state = "veryHidden"
    wb = openpyxl.Workbook()
    build(wb.active)
    wb.create_sheet("보임")["A1"] = 1
    path = tmp_path / "v.xlsx"
    wb.save(path)
    assert _open(path)[0].sheet("S").hidden is True


# ---------- 수식 ----------

def test_formula_without_saved_value(generated_dir):
    s = _open(generated_dir / "가상몰_변형.xlsx")[0].sheet("숨김행수식")
    g5 = s.cell(5, "G")
    assert (g5.value, g5.formula_no_value) == (None, True)
    assert g5.origin == "cell"          # 수식 칸 자체는 원문 칸
    assert s.cell(5, "F").formula_no_value is False


# ---------- 시트 종류·범위 ----------

def test_sheet_order_and_chart_sheet(generated_dir):
    wb, _ = _open(generated_dir / "가상몰_변형.xlsx")
    names = [s.name for s in wb.sheets]
    assert names.index("차트") == names.index("왼쪽덩어리") + 1      # 통합 문서 안 시트 순서
    chart = wb.sheet("차트")
    assert chart.is_chart is True and chart.max_row == 0 and chart.columns == []
    empty = wb.sheet("빈시트")
    assert empty.is_chart is False and empty.max_row == 0


def test_inflated_range_uses_real_values(generated_dir):
    s = _open(generated_dir / "가상몰_변형.xlsx")[0].sheet("부푼범위")
    assert s.max_row == 4                 # 500행 서식은 범위에 넣지 않음
    assert s.columns == ["A", "B", "C", "D", "E", "F", "G"]


def test_range_includes_merge_bounds(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["B2"] = "값"
        ws.merge_cells("B2:D5")           # 값은 B2뿐이지만 병합이 D5까지
    s = _open(_save(tmp_path, build))[0].sheet("S")
    assert (s.max_row, s.columns[-1]) == (5, "D")


def test_out_of_range_cell_is_empty(tmp_path):
    s = _open(_save(tmp_path, _merged_sheet))[0].sheet("S")
    c = s.cell(100, "Z")
    assert (c.value, c.origin) == (None, "empty")


# ---------- 칸 메모 ----------

def test_comments_not_read(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["A1"] = "값"
        ws["A1"].comment = Comment("메모 내용", "작성자")
        ws["B1"].comment = Comment("빈 칸 메모", "작성자")
    s = _open(_save(tmp_path, build))[0].sheet("S")
    assert s.cell(1, "A").value == "값"
    assert s.columns == ["A"]                                   # 메모만 있는 칸은 범위에 안 들어감
    assert not any(hasattr(s.cell(1, "A"), k) for k in ("comment", "note"))


# ---------- 파일 형식 ----------

@pytest.mark.parametrize("name", ["가상몰_옛형식.xls", "가상몰_가짜확장자.xlsx"])
def test_old_format_is_unsupported(generated_dir, name):
    wb, issues = _open(generated_dir / name)
    assert wb is None
    [issue] = issues.items
    assert issue.kind == "unsupported" and issue.file == name and issue.locations == ()


def test_csv_and_unknown_extension_unsupported(tmp_path):
    for name in ("a.csv", "b.txt"):
        path = tmp_path / name
        path.write_text("가,나\n", encoding="utf-8")
        wb, issues = _open(path)
        assert wb is None and [i.kind for i in issues] == ["unsupported"]
        assert name in issues.items[0].message


def test_xlsx_that_is_not_zip_is_unsupported(tmp_path):
    path = tmp_path / "깨짐.xlsx"
    path.write_bytes(b"not a zip file at all")
    wb, issues = _open(path)
    assert wb is None and [i.kind for i in issues] == ["read_failed"]   # 1.2: 처리 실패는 read_failed


def test_ole_header_with_xlsx_name_mentions_password(tmp_path):
    # 암호 걸린 .xlsx는 OLE 컨테이너로 저장된다(옛 형식과 같은 파일 앞부분). 실제 암호 파일 대신 앞부분이 같은 파일로 확인
    path = tmp_path / "암호.xlsx"
    path.write_bytes((generated_dir_path() / "가상몰_가짜확장자.xlsx").read_bytes())
    wb, issues = _open(path)
    assert wb is None and "암호" in issues.items[0].message


def generated_dir_path():
    return Path(__file__).parent / "fixtures" / "generated"


def test_lock_file_is_not_input(tmp_path):
    path = _save(tmp_path, _merged_sheet, name="~$t.xlsx")
    wb, issues = _open(path)
    assert wb is None and [i.kind for i in issues] == ["unsupported"]


def test_xlsm_supported(tmp_path):
    path = _save(tmp_path, _merged_sheet, name="m.xlsm")
    wb, issues = _open(path)
    assert wb is not None and len(issues) == 0 and wb.file == "m.xlsm"


def test_workbook_file_is_name_only(generated_dir):
    wb, _ = _open(generated_dir / "가상몰_변형.xlsx")
    assert wb.file == "가상몰_변형.xlsx"


def test_unknown_sheet_name_is_error(tmp_path):
    wb, _ = _open(_save(tmp_path, _merged_sheet))
    with pytest.raises(KeyError):
        wb.sheet("없음")


# ---------- openpyxl은 reader만 ----------

def test_only_reader_imports_openpyxl():
    users = []
    for path in sorted(SRC.glob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            names = [a.name for a in node.names] if isinstance(node, ast.Import) else \
                [node.module or ""] if isinstance(node, ast.ImportFrom) else []
            if any(n == "openpyxl" or n.startswith("openpyxl.") for n in names):
                users.append(path.name)
    assert sorted(set(users)) == ["reader.py"]


# ---------- 검토 반영: 큰 부푼 범위, 열기 실패, 병합 시작 칸 수식 ----------

def test_huge_inflated_range_reads_only_real_cells(tmp_path):
    import time as _time
    from openpyxl.styles import Font

    def build(ws):
        ws.title = "S"
        ws["A1"] = "머리글"
        ws["A2"] = "값"
        ws.cell(1048576, 30).font = Font(bold=True)   # 맨 끝 행·30열에 서식만 → max_row·max_column이 크게 부풂
    path = _save(tmp_path, build)
    started = _time.perf_counter()
    s = _open(path)[0].sheet("S")
    assert _time.perf_counter() - started < 10     # 부푼 범위 전체(약 3천만 칸)를 돌면 넘는다
    assert (s.max_row, s.columns) == (2, ["A"])


def test_permission_error_is_reported_not_raised(tmp_path, monkeypatch):
    path = _save(tmp_path, _merged_sheet)

    def deny(*a, **k):
        raise PermissionError(13, "Permission denied")
    monkeypatch.setattr(reader.openpyxl, "load_workbook", deny)
    wb, issues = _open(path)
    assert wb is None and [i.kind for i in issues] == ["read_failed"]   # 1.2: 처리 실패는 read_failed
    assert "처리 실패" in issues.items[0].message and "열 수 없습니다" in issues.items[0].message


def test_corrupt_zip_contents_reported_not_raised(tmp_path):
    import zipfile as _zip
    path = tmp_path / "손상.xlsx"
    with _zip.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", "<<깨진 XML")
    wb, issues = _open(path)
    assert wb is None and [i.kind for i in issues] == ["read_failed"]   # 1.2: 처리 실패는 read_failed
    assert "처리 실패" in issues.items[0].message


def test_unsupported_format_message_says_out_of_scope(generated_dir):
    _, issues = _open(generated_dir / "가상몰_옛형식.xls")
    assert "지원 범위 밖" in issues.items[0].message


def test_merged_cells_under_formula_start_without_value(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["B2"] = "=1+1"                  # openpyxl로 저장 → 계산값 없음
        ws.merge_cells("B2:B3")
    s = _open(_save(tmp_path, build))[0].sheet("S")
    b3 = s.cell(3, "B")
    assert (b3.value, b3.origin, b3.formula_no_value, b3.merge_start) == (None, "empty", True, (2, "B"))


# ---------- 값 처리 결정 ----------

def test_elapsed_time_becomes_total_hours_text(tmp_path):
    from datetime import timedelta

    def build(ws):
        ws.title = "S"
        ws["A1"] = timedelta(hours=26, minutes=30)
        ws["A1"].number_format = "[h]:mm"
        ws["B1"] = timedelta(minutes=5, seconds=7)
        ws["B1"].number_format = "[h]:mm:ss"
    s = _open(_save(tmp_path, build))[0].sheet("S")
    assert [s.cell(1, c).value for c in "AB"] == ["26:30:00", "0:05:07"]


def test_number_format_recorded_for_numbers(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["A1"] = 1
        ws["A1"].number_format = "000"
        ws["B1"] = 2                       # 일반 서식
        ws["C1"] = "글자"
    s = _open(_save(tmp_path, build))[0].sheet("S")
    a1 = s.cell(1, "A")
    assert (a1.value, a1.number_format) == (1, "000")      # 값은 저장 값 그대로, 서식은 기록만
    assert s.cell(1, "B").number_format is None and s.cell(1, "C").number_format is None


def test_formula_error_value_kept_and_marked(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["A1"] = "#N/A"                  # openpyxl은 오류 값으로 저장
        ws["B1"] = "보통 글자"
    s = _open(_save(tmp_path, build))[0].sheet("S")
    a1 = s.cell(1, "A")
    assert (a1.value, a1.error, a1.origin) == ("#N/A", True, "cell")
    assert s.cell(1, "B").error is False


def test_formula_error_marked_on_merge_restored(tmp_path):
    def build(ws):
        ws.title = "S"
        ws["A1"] = "#REF!"
        ws.merge_cells("A1:A2")
    a2 = _open(_save(tmp_path, build))[0].sheet("S").cell(2, "A")
    assert (a2.value, a2.origin, a2.error) == ("#REF!", "merge_restored", True)


def test_empty_string_is_empty_cell(tmp_path):
    import zipfile as _zip
    # openpyxl은 ""를 저장하지 않으므로, 저장된 파일 안의 공유 문자열에 빈 글자를 넣어 만든다
    def build(ws):
        ws.title = "S"
        ws["A1"] = "머리글"
        ws["A2"] = "EMPTYMARK"
        ws["A3"] = "  "                    # 공백만 있는 칸은 원문 값으로 남김
    path = _save(tmp_path, build)
    fixed = tmp_path / "e.xlsx"
    with _zip.ZipFile(path) as src, _zip.ZipFile(fixed, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith(".xml"):
                data = data.replace(b"EMPTYMARK", b"")
            dst.writestr(item, data)
    s = _open(fixed)[0].sheet("S")
    a2 = s.cell(2, "A")
    assert (a2.value, a2.origin) == (None, "empty")
    assert s.cell(3, "A").value == "  "


# ---------- 6c: 시트 이름만 읽기 ----------

def test_sheet_names_in_order(generated_dir):
    from si_prep.reader import sheet_names
    issues = Issues()
    names = sheet_names(generated_dir / "가상몰_변형.xlsx", issues)
    assert names[:2] == ["빈칸채움", "빈칸채움_끔"] and "차트" in names and "숨김" in names and len(issues) == 0


def test_sheet_names_unsupported_file(generated_dir):
    from si_prep.reader import sheet_names
    issues = Issues()
    assert sheet_names(generated_dir / "가상몰_옛형식.xls", issues) is None
    assert [i.kind for i in issues] == ["unsupported"]


def test_sheet_list_hidden_and_chart(generated_dir):
    from si_prep.reader import sheet_list
    info = {i.name: i for i in sheet_list(generated_dir / "가상몰_변형.xlsx", Issues())}
    assert (info["숨김"].hidden, info["숨김"].is_chart) == (True, False)
    assert (info["차트"].hidden, info["차트"].is_chart) == (False, True)
    assert (info["빈칸채움"].hidden, info["빈칸채움"].is_chart) == (False, False)
