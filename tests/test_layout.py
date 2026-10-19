"""layout ① 표 읽기 기본 — 시트 선택, 두 줄 머리글 합치기, 선택자 → 열 문자, header_not_found,
끝 표시, hidden_rows: skip. 표 배치 경고(표 여러 개·떨어진 덩어리·각주)는 9b.
"""

from pathlib import Path

import openpyxl
import pytest

from si_prep.config import load_profile, sheet_profile
from si_prep.issues import Issues
from si_prep.layout import build_table, select_sheets
from si_prep.reader import open_workbook
from conftest import make_profile, write_book

PROFILES = Path(__file__).parent / "fixtures" / "configs" / "shop" / "profiles"
FO = "가상몰_기능정의서_v1.2.xlsx"
VAR = "가상몰_변형.xlsx"


def _wb(generated_dir, name):
    return open_workbook(generated_dir / name, Issues())


def _table(generated_dir, file, profile_name, sheet, required=()):
    profile = load_profile(PROFILES / f"{profile_name}.yaml")
    grid = _wb(generated_dir, file).sheet(sheet)
    issues = Issues()
    return build_table(grid, sheet_profile(profile, sheet), file, issues, required_roles=required), issues


def _profile(**over):
    return make_profile({"hierarchy": [{"header_contains": "분류"}]}, **over)


def _grid(tmp_path, rows, name="S", merges=(), hidden_rows=()):
    path = write_book(tmp_path / "t.xlsx", rows, merges, sheet=name, hidden_rows=hidden_rows)
    return open_workbook(path, Issues()).sheet(name)


# ---------- 시트 선택 ----------

def test_select_sheets_excludes_read_false_in_workbook_order(generated_dir):
    wb = _wb(generated_dir, FO)
    issues = Issues()
    chosen = select_sheets(wb, load_profile(PROFILES / "fo.yaml"), issues)
    assert [g.name for g in chosen] == [s.name for s in wb.sheets if s.name not in ("표지", "개정이력")]
    assert len(issues) == 0                                   # 뺌은 프로필 선택이라 경고 없음


def test_select_sheets_hidden_and_chart_skipped(generated_dir):
    wb = _wb(generated_dir, VAR)
    issues = Issues()
    names = [g.name for g in select_sheets(wb, load_profile(PROFILES / "variation.yaml"), issues)]
    assert "숨김" not in names and "차트" not in names and "빈시트" in names
    skipped = {(i.file, i.sheet): i for i in issues.by_kind("sheet_skipped")}
    assert set(skipped) == {(VAR, "숨김"), (VAR, "차트")}
    assert "chart_sheet" in skipped[(VAR, "차트")].message and skipped[(VAR, "숨김")].locations == ()


def test_select_sheets_pattern_and_new_sheet(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "표지1"
    for name in ("표지2", "본문", "새시트"):
        wb.create_sheet(name)["A1"] = "x"
    path = tmp_path / "p.xlsx"
    wb.save(path)
    profile = _profile(sheets=[{"name": "표지*", "read": False, "reason": "메타"}, {"name": "본문"}])
    chosen = select_sheets(open_workbook(path, Issues()), profile, Issues())
    assert [g.name for g in chosen] == ["본문", "새시트"]     # 목록에 없는 새 시트는 기본값으로 읽음


def test_select_sheets_exact_name_beats_pattern(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "표지 요약"
    wb.active["A1"] = "x"
    path = tmp_path / "e.xlsx"
    wb.save(path)
    profile = _profile(sheets=[{"name": "표지*", "read": False, "reason": "메타"}, {"name": "표지 요약"}])
    assert [g.name for g in select_sheets(open_workbook(path, Issues()), profile, Issues())] == ["표지 요약"]


# ---------- 머리글 합치기·역할 ----------

def test_two_row_header_joined_once_for_merges(generated_dir):
    table, issues = _table(generated_dir, FO, "fo", "회원")
    assert table.headers["B"] == "분류 / 1단계" and table.headers["D"] == "분류 / 3단계"
    assert table.headers["E"] == "기능명" and table.headers["F"] == "설명"     # 세로 병합은 한 번만
    assert len(issues) == 0


def test_selectors_to_columns(generated_dir):
    table, _ = _table(generated_dir, FO, "fo", "회원")
    assert table.roles == {"hierarchy": ["B", "C", "D"], "title": ["E"], "description": ["F"]}
    assert (table.file, table.sheet, table.header_rows) == (FO, "회원", [4, 5])


def test_override_columns_used(generated_dir):
    table, _ = _table(generated_dir, FO, "fo", "고객센터")
    assert table.roles["hierarchy"] == ["C", "D", "E", "F"] and table.roles["title"] == ["G"]
    assert table.headers["B"] == "NO." and table.headers["I"] == "비고"       # 역할 없는 열도 머리글은 있음


@pytest.mark.parametrize("sheet,cols", [("분류패턴_2단", ["B", "C"]), ("분류패턴_3단", ["B", "C", "D"])])
def test_header_pattern_all_expands(generated_dir, sheet, cols):
    table, issues = _table(generated_dir, VAR, "variation", sheet)
    assert table.roles["hierarchy"] == cols and len(issues) == 0


def test_header_pattern_expands_in_place_with_other_selectors(tmp_path):
    grid = _grid(tmp_path, [["화면", "분류1", "분류2", "케이스"], ["가", "나", "다", "라"]])
    p = _profile(columns={"hierarchy": [{"header_contains": "화면"}, {"header_pattern": r"^분류\d$", "all": True}],
                          "title": {"header_contains": "케이스"}})
    table = build_table(grid, p, "t.xlsx", Issues())
    assert table.roles["hierarchy"] == ["A", "B", "C"]


def test_col_selector(tmp_path):
    grid = _grid(tmp_path, [["세부항목", "세부항목", "내용"], ["가", "나", "다"]], merges=["A1:B1"])
    p = _profile(columns={"hierarchy": [{"col": "A"}, {"col": "B"}], "description": {"header_contains": "내용"}})
    table = build_table(grid, p, "t.xlsx", Issues())
    assert table.roles == {"hierarchy": ["A", "B"], "description": ["C"]}
    assert table.headers == {"A": "세부항목", "B": "세부항목", "C": "내용"}


def test_exclude_from_text_columns(tmp_path):
    grid = _grid(tmp_path, [["분류", "결과(가)", "결과(나)"], ["가", "P", "F"]])
    p = _profile(columns={"hierarchy": [{"header_contains": "분류"}],
                          "exclude_from_text": [{"header_contains": "결과(가)"}, {"header_contains": "결과(나)"}]})
    assert build_table(grid, p, "t.xlsx", Issues()).roles["exclude_from_text"] == ["B", "C"]


# ---------- header_not_found ----------

def test_single_value_role_matching_two_columns(generated_dir):
    table, issues = _table(generated_dir, VAR, "variation", "역할겹침")
    assert table is None
    [w] = issues.items
    assert w.kind == "header_not_found" and (w.file, w.sheet) == (VAR, "역할겹침")
    assert [(l.rows.start, l.rows.end, l.cols) for l in w.locations] == [(1, 1, ["E", "F"])]
    assert "description" in w.message and "E" in w.message and "F" in w.message


def test_missing_hierarchy_selector(generated_dir):
    table, issues = _table(generated_dir, VAR, "variation", "분류열부족")
    assert table is None
    [w] = issues.items
    assert w.kind == "header_not_found" and "분류3" in w.message
    assert [(l.rows.start, l.rows.end, l.cols) for l in w.locations] == [(1, 1, None)]


def test_empty_header_rows(generated_dir):
    table, issues = _table(generated_dir, VAR, "variation", "빈시트")
    assert table is None and [i.kind for i in issues] == ["header_not_found"]
    assert "비어" in issues.items[0].message


def test_header_pattern_no_match(tmp_path):
    grid = _grid(tmp_path, [["구분", "설명"], ["가", "나"]])
    p = _profile(columns={"hierarchy": [{"header_pattern": r"^분류\d$", "all": True}]})
    issues = Issues()
    assert build_table(grid, p, "t.xlsx", issues) is None and [i.kind for i in issues] == ["header_not_found"]


def test_hierarchy_not_required_for_none(tmp_path):
    grid = _grid(tmp_path, [["이름", "설명"], ["가", "나"]])
    p = _profile(columns={"title": {"header_contains": "이름"}}, hierarchy={"method": "none"})
    table = build_table(grid, p, "t.xlsx", Issues())
    assert table.roles == {"title": ["A"]}


def test_optional_role_missing_is_not_error(tmp_path):
    grid = _grid(tmp_path, [["분류", "설명"], ["가", "나"]])
    p = _profile(columns={"hierarchy": [{"header_contains": "분류"}], "title": {"header_contains": "기능명"}})
    issues = Issues()
    table = build_table(grid, p, "t.xlsx", issues)
    assert table is not None and "title" not in table.roles and len(issues) == 0
    assert table.matches["title"] == []                       # 설정 검사 출력용: 선택자마다 맞은 열


def test_continuation_requires_id_column(tmp_path):
    grid = _grid(tmp_path, [["분류", "설명"], ["가", "나"]])
    p = _profile(item_boundary={"method": "continuation"})
    issues = Issues()
    assert build_table(grid, p, "t.xlsx", issues) is None and "id" in issues.items[0].message


def test_required_roles_from_caller(tmp_path):
    # id_ref 대상 쪽이면 run이 id를 필수로 넘긴다
    grid = _grid(tmp_path, [["분류", "설명"], ["가", "나"]])
    issues = Issues()
    assert build_table(grid, _profile(), "t.xlsx", issues, required_roles=("id",)) is None
    assert [i.kind for i in issues] == ["header_not_found"] and "id" in issues.items[0].message
    assert build_table(grid, _profile(), "t.xlsx", Issues()) is not None


# ---------- 데이터 행 ----------

def test_data_rows_stop_at_marker(generated_dir):
    table, _ = _table(generated_dir, FO, "fo", "주문")
    assert table.data_rows == [6, 7, 8]                         # 9행 「※ 참고」부터 무시


def test_data_rows_until_real_end(generated_dir):
    table, _ = _table(generated_dir, VAR, "variation", "부푼범위")
    assert table.data_rows == [2, 3, 4]


def test_hidden_rows_skip(generated_dir):
    table, issues = _table(generated_dir, VAR, "variation", "숨김건너뜀")
    assert table.data_rows == [4, 5]
    [w] = issues.items
    assert w.kind == "hidden_content" and "건너" in w.message
    assert [(l.rows.start, l.rows.end) for l in w.locations] == [(2, 3)]


def test_hidden_rows_include_keeps_rows(generated_dir):
    table, issues = _table(generated_dir, VAR, "variation", "숨김행수식")
    assert table.data_rows == [2, 3, 4, 5]                      # 숨김 행도 처리(include의 hidden_content는 9b)


def test_matches_report_columns_and_headers(generated_dir):
    table, _ = _table(generated_dir, FO, "fo", "회원")
    m = table.matches
    assert m["title"] == [("E", "기능명")]
    assert m["hierarchy"] == [("B", "분류 / 1단계"), ("C", "분류 / 2단계"), ("D", "분류 / 3단계")]


# ---------- 검토 반영 ----------

def test_col_selector_outside_sheet_is_not_found(tmp_path):
    grid = _grid(tmp_path, [["분류", "설명"], ["가", "나"]])
    p = _profile(columns={"hierarchy": [{"col": "Z"}]})
    issues = Issues()
    assert build_table(grid, p, "t.xlsx", issues) is None and "col" in issues.items[0].message


def test_hierarchy_duplicate_columns_once(tmp_path):
    grid = _grid(tmp_path, [["분류1", "분류2"], ["가", "나"]])
    p = _profile(columns={"hierarchy": [{"header_contains": "분류1"}, {"header_pattern": r"^분류\d$", "all": True}]})
    assert build_table(grid, p, "t.xlsx", Issues()).roles["hierarchy"] == ["A", "B"]


def test_two_hierarchy_overlaps_both_reported(tmp_path):
    grid = _grid(tmp_path, [["분류 가", "분류 가2", "구분 나", "구분 나2"], ["1", "2", "3", "4"]])
    p = _profile(columns={"hierarchy": [{"header_contains": "분류 가"}, {"header_contains": "구분 나"}]})
    issues = Issues()
    assert build_table(grid, p, "t.xlsx", issues) is None
    assert issues.items[0].locations[0].cols == ["A", "B", "C", "D"]


def test_skip_message_separates_outline_rows(generated_dir):
    profile = load_profile(PROFILES / "variation.yaml")
    grid = _wb(generated_dir, VAR).sheet("숨김행수식")
    issues = Issues()
    sp = sheet_profile(profile, "숨김행수식").model_copy(update={"hidden_rows": "skip"})
    table = build_table(grid, sp, VAR, issues)
    assert table.data_rows == [2, 5]
    [w] = issues.items
    assert "그룹 접기 1줄" in w.message and "그냥 숨김 1줄" in w.message


def test_whitespace_only_header_is_empty(tmp_path):
    grid = _grid(tmp_path, [["  ", "분류"], ["가", "나"]])
    table = build_table(grid, _profile(), "t.xlsx", Issues())
    assert table.headers == {"B": "분류"}


def test_selectors_recorded_per_column(generated_dir):
    table, _ = _table(generated_dir, FO, "fo", "회원")
    assert [c for c, _ in table.selectors["hierarchy"]] == ["B", "C", "D"]
    assert table.selectors["title"][0][1].header_contains == "기능명"


def test_header_parts_mark_wide_merges(generated_dir):
    # 2칸 이상 가로 병합된 머리글 칸(묶음 제목)인지 칸마다 기록(설정 검사 출력에서 보일지 정하는 데 씀)
    table, _ = _table(generated_dir, FO, "fo", "회원")
    assert table.header_parts["B"] == [("분류", True), ("1단계", False)]
    assert table.header_parts["E"] == [("기능명", False)]                  # 세로 병합(E4:E5)은 가로 병합 아님
