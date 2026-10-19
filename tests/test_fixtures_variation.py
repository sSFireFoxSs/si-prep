"""항목 2-2: 변형 파일(표 모양·경고 시나리오)과 미지원 형식 파일의 구조를 확인한다.

시나리오는 docs/DESIGN.md 8.1. 기대값은 생성기가 넣은 칸으로만 본다.
"""

import zipfile
from pathlib import Path

import openpyxl
import yaml

VARIATION = "가상몰_변형.xlsx"
OLD_XLS = "가상몰_옛형식.xls"
FAKE_XLSX = "가상몰_가짜확장자.xlsx"
OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
CONFIG_DIR = Path(__file__).parent / "fixtures" / "configs" / "shop"

SHEETS = [
    "빈칸채움", "빈칸채움_끔", "한칸경로", "번호", "이어지는행", "오른쪽메모", "역할겹침", "숨김",
    "표두개", "들여쓰기", "숨김행수식", "숨김건너뜀", "내용병합", "분류패턴_2단", "분류패턴_3단",
    "분류열부족", "각주", "부푼범위", "표안빈열", "대각선표", "왼쪽덩어리", "차트", "빈시트", "표여러개",
]


def _wb(generated_dir, **kw):
    return openpyxl.load_workbook(generated_dir / VARIATION, **kw)


def _values(ws):
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}


def _row_empty(ws, r):
    return all(ws.cell(r, c).value is None for c in range(1, ws.max_column + 1))


def _col_empty(ws, col):
    return all(ws.cell(r, col).value is None for r in range(1, ws.max_row + 1))


def _headers(ws, row=1):
    return [c.value for c in ws[row] if c.value is not None]


# ---------- 시트 목록 ----------

def test_variation_sheet_list(generated_dir):
    wb = _wb(generated_dir)
    assert wb.sheetnames == SHEETS
    assert [c.title for c in wb.chartsheets] == ["차트"]  # D8: 칸 없는 차트 시트


# ---------- 분류 표현 ----------

def test_fill_down_on_off_same_source(generated_dir):
    wb = _wb(generated_dir)
    on, off = wb["빈칸채움"], wb["빈칸채움_끔"]
    assert _values(on) == _values(off)  # 켬/끔 비교용이라 원문이 같다
    assert _headers(on) == ["분류1", "분류2", "분류3", "기능명", "설명"]
    # docs/DESIGN.md 8.1 fill_down 케이스: 행 1~5 = 엑셀 2~6행, 분류 칸(B·C·D) 모양
    shape = [tuple(on.cell(r, c).value is not None for c in (2, 3, 4)) for r in range(2, 7)]
    assert shape == [
        (True, True, False),    # A / B / —
        (False, False, True),   # — / — / C
        (False, True, False),   # — / D / —
        (False, True, False),   # — / E / —
        (False, False, False),  # — / — / —  (분류 칸 전부 빈 행)
    ]
    assert not on.merged_cells.ranges  # 병합 없음


def test_split_cell_and_numbering(generated_dir):
    wb = _wb(generated_dir)
    assert wb["한칸경로"]["B2"].value.split(" > ") == ["회원", "로그인", "아이디 찾기"]
    numbers = [wb["번호"].cell(r, 2).value for r in range(2, 7)]
    assert numbers == ["1", "1.1", "1.1.1", "2", "2.3.1"]  # 2.3이 없음


def test_indent_sheet(generated_dir):
    ws = _wb(generated_dir)["들여쓰기"]
    assert [ws.cell(r, 2).alignment.indent for r in range(2, 5)] == [0, 1, 2]


def test_pattern_and_missing_hierarchy_headers(generated_dir):
    wb = _wb(generated_dir)
    assert _headers(wb["분류패턴_2단"]) == ["분류1", "분류2", "기능명", "설명"]
    assert _headers(wb["분류패턴_3단"]) == ["분류1", "분류2", "분류3", "기능명", "설명"]
    assert "분류3" not in _headers(wb["분류열부족"])  # 선택자 3개 중 하나가 없음


# ---------- 항목 경계·역할 ----------

def test_continuation_rows(generated_dir):
    ws = _wb(generated_dir)["이어지는행"]
    assert ws["B2"].value == "V-001"
    assert ws["B3"].value is None and ws["B4"].value is None  # ID 빈 다음 2행
    assert ws["E3"].value is None and ws["F3"].value  # 중간 행의 한 열은 빈칸
    assert ws["E4"].value and ws["F4"].value


def test_role_overlap(generated_dir):
    ws = _wb(generated_dir)["역할겹침"]
    assert sum("설명" in h for h in _headers(ws)) == 2  # description 선택자에 열 2개가 걸림


def test_content_merge(generated_dir):
    ws = _wb(generated_dir)["내용병합"]
    assert "F2:F4" in {str(r) for r in ws.merged_cells.ranges}
    assert ws["F2"].value and ws["F3"].value is None and ws["F4"].value is None
    assert ws["E2"].value and ws["E3"].value and ws["E4"].value  # 세 행 모두 각자 항목


# ---------- 숨김·수식 ----------

def test_hidden_sheet(generated_dir):
    assert _wb(generated_dir)["숨김"].sheet_state == "hidden"


def test_hidden_rows_plain_and_grouped_and_formula(generated_dir):
    ws = _wb(generated_dir)["숨김행수식"]
    plain, grouped = ws.row_dimensions[3], ws.row_dimensions[4]
    assert plain.hidden and plain.outline_level == 0      # 그냥 숨긴 행
    assert grouped.hidden and grouped.outline_level == 1  # 그룹 접기
    assert ws["G5"].data_type == "f"
    assert _wb(generated_dir, data_only=True)["숨김행수식"]["G5"].value is None  # 계산값 없음


def test_hidden_rows_skip_sheet(generated_dir):
    ws = _wb(generated_dir)["숨김건너뜀"]
    hidden = sorted(r for r, d in ws.row_dimensions.items() if d.hidden)
    assert hidden == [2, 3]  # 데이터 앞쪽 연속 숨김 행 2개
    assert ws["E4"].value and ws["E5"].value


# ---------- 표 배치 (R6 포함) ----------

def test_right_memo(generated_dir):
    ws = _wb(generated_dir)["오른쪽메모"]
    assert _headers(ws)[-1] == "설명" and ws["F1"].value == "설명"
    assert _col_empty(ws, 7)  # G열 비움
    assert ws["H2"].value and ws["H3"].value


def test_two_tables(generated_dir):
    ws = _wb(generated_dir)["표두개"]
    assert all(_row_empty(ws, r) for r in (4, 5, 6))  # 빈 행 3줄
    assert ws["B7"].value and ws["B8"].value


def test_footnote(generated_dir):
    ws = _wb(generated_dir)["각주"]
    assert _row_empty(ws, 4)
    row5 = [c.value for c in ws[5] if c.value is not None]
    assert len(row5) == 1 and not str(row5[0]).startswith("※")  # 끝 표시 글자 없음
    # 각주는 역할 없는 번호 열(B)에만 있고 역할 열(분류·기능명·설명)은 빈칸
    assert ws["B1"].value == "번호" and ws["B5"].value == row5[0]


def test_inflated_max_row(generated_dir):
    ws = _wb(generated_dir)["부푼범위"]
    assert ws.max_row >= 500  # D1: 서식만 있는 먼 칸
    assert max(c.row for row in ws.iter_rows() for c in row if c.value is not None) == 4


def test_empty_column_inside_header(generated_dir):
    ws = _wb(generated_dir)["표안빈열"]
    assert ws["E1"].value == "메모"
    assert sum(ws.cell(r, 5).value is not None for r in range(2, 5)) == 1  # D3: 값이 거의 없음


def test_diagonal_table(generated_dir):
    ws = _wb(generated_dir)["대각선표"]
    assert all(_row_empty(ws, r) for r in (4, 5, 6))
    second = [c for row in ws.iter_rows(min_row=7) for c in row if c.value is not None]
    assert second and min(c.column for c in second) > 6  # 본 표(B~F)의 오른쪽 아래


def test_left_block(generated_dir):
    ws = _wb(generated_dir)["왼쪽덩어리"]
    assert ws["D1"].value == "분류1"  # 헤더는 D부터
    assert _col_empty(ws, 3)  # C열 비움
    assert ws["B2"].value and ws["B3"].value


def test_empty_sheet(generated_dir):
    assert not _values(_wb(generated_dir)["빈시트"])  # U5


# ---------- 미지원 형식 파일 ----------

def test_unsupported_format_files(generated_dir):
    for name in (OLD_XLS, FAKE_XLSX):
        path = generated_dir / name
        assert path.read_bytes()[:8] == OLE_SIGNATURE, name
        assert not zipfile.is_zipfile(path), name


# ---------- 설정 ----------

def test_variation_profile_overrides_match_sheets(generated_dir):
    profile = yaml.safe_load((CONFIG_DIR / "profiles" / "variation.yaml").read_text(encoding="utf-8"))
    names = set(_wb(generated_dir).sheetnames)
    for override in profile["sheet_overrides"]:
        assert set(override["sheets"]) <= names, override["sheets"]
