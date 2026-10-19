"""합성 엑셀 생성기 (가상 쇼핑몰).

실제 문서의 내용·이름을 쓰지 않는다. 시나리오는 docs/DESIGN.md 8.1.
항목 2-1: 주 문서 4종(요구사항·기능·정책·테스트).
항목 2-2: 변형 파일(표 모양·경고 시나리오)과 미지원 형식 파일.
항목 2-4: 관리자 기능정의서(bo), 테스트 BO 시트·담당자별 결과 열, 정책 서술형·목록형 시트.
"""

from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font

# 문서 속성의 만든 시각은 고정하지만, 저장 시각(modified)은 openpyxl이 저장할 때마다 바꾼다.
# 그래서 파일 바이트는 매번 달라진다. "같음"은 구조·칸 값 기준으로 본다(tests/test_fixtures.py 결정성 테스트)
FIXED_TIME = datetime(2026, 10, 1, 9, 0, 0)


def _new_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "make_fixtures"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    return wb


def _put(ws, cells):
    """{"B2": 값, ...}를 칸에 쓴다."""
    for ref, value in cells.items():
        ws[ref] = value


def _rows(ws, first_row, first_col, rows):
    """first_row부터 행 목록을 쓴다. None인 칸은 비워 둔다."""
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            if value is not None:
                ws.cell(first_row + i, first_col + j, value)


# ---------- 요구사항정의서 (req) ----------

def make_req(path):
    wb = _new_workbook()
    ws = wb.create_sheet("요구사항")
    _rows(ws, 1, 2, [["대분류", "중분류", "요구사항 ID", "요구사항명", "설명"]])
    _rows(ws, 2, 2, [
        ["고객", "로그인", "REQ-001", "로그인 기능", "아이디와 비밀번호로 로그인"],
        [None, "문의", "REQ-002", "1:1 문의 기능", "문의 등록과 답변 확인"],
        ["주문", "결제", "REQ-003", "결제 기능", "카드와 무통장 입금 지원"],
        ["주문", "배송", "REQ-003", "배송지 변경 기능", "주문 후 배송지 1회 변경"],  # 같은 ID 두 번
    ])
    ws.merge_cells("B2:B3")  # 대분류 "고객" 병합 2행
    wb.save(path)


# ---------- 기능정의서 (fo) ----------

def _fo_cover_and_history(wb):
    cover = wb.create_sheet("표지")
    _put(cover, {"B2": "가상몰 기능정의서", "B4": "문서명", "C4": "기능정의서", "B5": "버전", "C5": "v1.3"})
    cover.merge_cells("B2:C2")
    history = wb.create_sheet("개정이력")
    _rows(history, 2, 2, [["버전", "일자", "변경 사항", "작성"], ["v1.0", "2026-09-01", "최초 작성", "기획팀"]])


def _fo_menu_header(ws):
    """두 줄 헤더(4·5행): 분류(1~3단계) / 기능명 / 설명"""
    _put(ws, {"B4": "분류", "B5": "1단계", "C5": "2단계", "D5": "3단계", "E4": "기능명", "F4": "설명"})
    for ref in ("B4:D4", "E4:E5", "F4:F5"):
        ws.merge_cells(ref)


def _fo_real_layout_sheet(wb):
    """흔한 SI 양식 구조를 따른 메뉴 시트(서식 없음): 제목·업데이트 행, NO. 열, 4칸 병합 분류(4단계),
    12행 세로 병합, 「- 」 목록이 있는 여러 줄 설명, 대부분 빈 비고 열."""
    ws = wb.create_sheet("고객센터")
    _put(ws, {"B1": "고객센터 기능정의", "B2": "문서 업데이트: 2026-09-28",
              "B4": "NO.", "C4": "분류", "C5": "1단계", "D5": "2단계", "E5": "3단계", "F5": "4단계",
              "G4": "기능명", "H4": "설명", "I4": "비고"})
    for ref in ("B4:B5", "C4:F4", "G4:G5", "H4:H5", "I4:I5"):
        ws.merge_cells(ref)
    _rows(ws, 6, 2, [
        [1, "문의", "문의 작성", "문의 종류", "주문 문의", "주문 문의 선택",
         "주문 번호를 함께 고른다\n- 최근 3개월 주문만\n- 취소 주문 포함", None],
        [2, None, None, None, "배송 문의", "배송 문의 선택", "배송 단계를 함께 보여 준다", None],
        [3, None, None, None, "기타 문의", "기타 문의 선택", "자유 입력", None],
        [4, None, None, "본문 입력", None, "본문 입력란", "1,000자까지 입력\n- 줄바꿈 허용", None],
        [5, None, None, "사진 첨부", None, "사진 첨부 버튼", "사진을 붙인다", "최대 3장"],
        [6, None, "문의 내역", "목록", None, "문의 목록 화면", "최근 문의부터 표시", None],
        [7, None, None, "상세", "답변 보기", "답변 보기 화면",
         "답변과 첨부를 함께 보여 준다\n- 답변 대기 표시\n- 답변 완료 표시", None],
        [8, None, None, None, "추가 문의", "추가 문의 버튼", "같은 문의에 이어 묻기", None],
        [9, None, None, "삭제", None, "문의 삭제 버튼", "답변 전에만 삭제", "확인 필요"],
        [10, None, "답변 알림", "앱 알림", None, "앱 알림 설정", "답변 등록 시 푸시", None],
        [11, None, None, "문자 알림", None, "문자 알림 설정", "답변 등록 시 문자", None],
        [12, None, None, "메일 알림", None, "메일 알림 설정", "답변 등록 시 메일", None],
    ])
    for ref in ("C6:C17", "D6:D10", "D11:D14", "D15:D17", "E6:E8", "E12:E13"):
        ws.merge_cells(ref)


def make_fo(path):
    wb = _new_workbook()
    _fo_cover_and_history(wb)

    member = wb.create_sheet("회원")
    _fo_menu_header(member)
    _rows(member, 6, 2, [
        ["로그인", "일반 로그인", "아이디 입력", "아이디 입력란", "영문·숫자 4~20자"],
        [None, None, "비밀번호 입력", "비밀번호 입력란", "8자 이상"],
        [None, "간편 로그인", "소셜 계정 선택", "소셜 로그인 버튼", None],    # 설명 빈칸 → not_written
        ["회원가입", "약관 동의", None, "약관 동의 화면", "필수 약관 3종 동의"],  # 3단계 빈칸 → 2단계에서 끝
        ["회원탈퇴", None, "탈퇴 사유", "탈퇴 사유 선택", "사유 5종 중 선택"],    # 2단계 빈칸(병합 아님)
    ])
    member.merge_cells("B6:B8")
    member.merge_cells("C6:C7")

    order = wb.create_sheet("주문")
    _fo_menu_header(order)
    _rows(order, 6, 2, [
        ["결제", "결제 수단", "카드 결제", "카드 결제 버튼", "국내 카드만 지원"],
        [None, None, "무통장 입금", "입금 계좌 안내", "입금 기한 3일"],
        ["배송", "배송지변경", None, "배송지 변경 화면", "주문 후 1회 변경 가능"],  # 테스트 쪽은 "배송지 변경"
        ["※ 참고: 배송 정책은 정책정의서를 따른다"],                              # 끝 표시 행
    ])
    order.merge_cells("B6:B7")
    order.merge_cells("C6:C7")

    _fo_real_layout_sheet(wb)

    overview = wb.create_sheet("요약")
    _put(overview, {"B4": "메뉴", "C4": "분류", "C5": "1단계", "D5": "2단계", "E5": "3단계",
                    "F4": "기능명", "G4": "설명"})
    for ref in ("B4:B5", "C4:E4", "F4:F5", "G4:G5"):
        overview.merge_cells(ref)
    _rows(overview, 6, 2, [
        ["회원", "로그인", "일반 로그인", "아이디 입력", "아이디 입력란", "영문·숫자 4~20자"],  # 메뉴 시트와 설명 같음
        ["주문", "결제", "결제 수단", "카드 결제", "카드 결제 버튼", "해외 카드도 지원"],      # 설명 다름
    ])
    wb.save(path)


# ---------- 관리자 기능정의서 (bo) ----------

def make_bo(path):
    """흔한 SI 양식처럼 1단계 열에 메뉴 이름이 있고, 시트 이름은 경로에 쓰지 않는다."""
    wb = _new_workbook()
    sheets = {
        "회원관리": [
            ["회원관리", "회원 목록", "검색", "회원 검색", "이름·가입일로 검색"],
            [None, None, "엑셀 내려받기", "내려받기 버튼", "검색 결과만 내려받기"],
            [None, "회원 상세", None, "회원 상세 화면", "등급과 주문 이력 표시"],
        ],
        "주문관리": [
            ["주문관리", "주문 목록", "상태 변경", "상태 변경 버튼", "결제 완료에서 배송 준비로"],
            [None, None, "송장 입력", "송장 입력란", "택배사 선택 후 입력"],  # 테스트 쪽은 "송장입력"
            [None, "환불 처리", None, "환불 처리 화면", "전액 환불만 지원"],
        ],
    }
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        _fo_menu_header(ws)
        _rows(ws, 6, 2, rows)
        ws.merge_cells("B6:B8")
        ws.merge_cells("C6:C7")
    wb.save(path)


# ---------- 정책정의서 (policy) ----------

def make_policy(path):
    wb = _new_workbook()
    sheets = {
        "주문 - 결제": [["결제 한도", "1회 결제 최대 200만 원"], ["결제 취소", "배송 시작 전까지 취소 가능"]],
        "회원 - 로그인": [["로그인 실패 제한", "5회 실패 시 10분 잠금"]],
        "쿠폰 - 발급": [["쿠폰 발급 조건", "첫 구매 고객에게 1장"]],  # 기능정의서에 없는 경로
    }
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        _rows(ws, 1, 2, [["정책명", "정책 설명"]] + rows)

    # 서술형: 2단 구분(머리글 B1:C1 병합) + 정책 내용 + 비고. 프로필 시트별 덮어쓰기로 처리
    narrative = wb.create_sheet("주문 - 배송")
    _rows(narrative, 1, 2, [
        ["세부 구분", None, "정책 내용", "비고"],
        ["배송비", "기본", "3만 원 이상 주문은 무료", None],
        [None, "도서 산간", "추가 배송비 부과", "금액은 지역별로 다름"],
        ["배송 기간", None, "결제 후 2~3일 안에 출고", None],
    ])
    narrative.merge_cells("B1:C1")
    narrative.merge_cells("B2:B3")

    # 목록형: 머리글 모양이 다른 표. 시트 이름에 " - "가 없어 연결 대상에서 뺀다(policy_fo sheets 패턴)
    listing = wb.create_sheet("결제 수단 목록")
    _rows(listing, 1, 2, [
        ["수단 코드", "수단 이름", "수수료율", "사용 여부"],
        ["PAY-01", "신용카드", "2.5%", "사용"],
        ["PAY-02", "무통장 입금", "0%", "사용"],
        ["PAY-03", "휴대폰 결제", "3.0%", "중지"],
    ])
    wb.save(path)


# ---------- 테스트케이스 (test) ----------

TEST_HEADER = ["케이스 ID", "요구사항 ID", "메뉴", "1단계", "2단계", "3단계", "테스트명", "기대결과", "수행자",
               "결과(검수자1)", "결과(검수자2)", "결과(검수자3)"]  # 결과 열: 담당자별 수행 결과(청크 본문 제외)


def make_test(path):
    wb = _new_workbook()
    fo = wb.create_sheet("테스트케이스")  # FO 대상
    _rows(fo, 1, 2, [
        TEST_HEADER,
        ["TC-001", "REQ-001", "회원", "로그인", "일반 로그인", "아이디 입력", "아이디 입력 확인", "입력한 아이디가 표시됨", "검수자1",
         "Pass", "Pass", None],
        ["TC-002", "REQ-001, REQ-002", "회원", "로그인", "일반 로그인", "비밀번호 입력", "비밀번호 가림 확인", "입력 글자가 가려짐", "검수자1",
         "Pass", None, None],
        ["TC-003", "REQ-999", "주문", "결제", "결제 수단", "카드 결제", "카드 결제 확인", "결제 완료 화면 표시", "검수자2",
         "Fail", "Pass", "Pass"],
        ["TC-004", None, "주문", "배송", "배송지 변경", None, "배송지 변경 확인", "변경된 주소 저장", "검수자2",
         None, None, None],
        ["TC-005", "요구사항 미정", "쿠폰", "쿠폰함", "쿠폰 목록", None, "쿠폰 목록 확인", "보유 쿠폰 표시", "검수자1",
         "Pass", None, None],
    ])
    bo = wb.create_sheet("BO테스트케이스")  # BO 대상 (머리글은 FO와 같음)
    _rows(bo, 1, 2, [
        TEST_HEADER,
        ["TC-101", None, "회원관리", "회원 목록", "검색", None, "회원 검색 확인", "검색 결과 표시", "검수자1",
         "Pass", "Pass", None],
        ["TC-102", None, "주문관리", "주문 목록", "상태 변경", None, "상태 변경 확인", "변경된 상태 저장", "검수자2",
         "Pass", "Fail", None],
        ["TC-103", None, "주문관리", "주문 목록", "송장입력", None, "송장 입력 확인", "송장 번호 저장", "검수자2",
         "Fail", None, None],
    ])
    wb.save(path)


# ---------- 변형 파일 (variation) — 표 모양·경고 시나리오 ----------

BASE_HEADER = ["분류1", "분류2", "분류3", "기능명", "설명"]  # B~F열


def _base_sheet(wb, name, rows, header=BASE_HEADER):
    """1행 헤더 + 2행부터 데이터인 기본 모양 시트"""
    ws = wb.create_sheet(name)
    _rows(ws, 1, 2, [header] + rows)
    return ws


def _fill_down_rows():
    # docs/DESIGN.md 8.1 fill_down 케이스 5행 (A=회원, B=로그인, C=아이디 찾기, D=회원가입, E=회원탈퇴)
    return [
        ["회원", "로그인", None, "로그인 화면", "아이디와 비밀번호 입력"],
        [None, None, "아이디 찾기", "아이디 찾기 화면", "휴대폰 인증으로 찾기"],
        [None, "회원가입", None, "가입 화면", "약관 동의 후 가입"],
        [None, "회원탈퇴", None, "탈퇴 화면", "탈퇴 사유 선택"],
        [None, None, None, "탈퇴 완료 안내", "완료 문구 표시"],
    ]


def make_variation(path):
    wb = _new_workbook()

    _base_sheet(wb, "빈칸채움", _fill_down_rows())       # 프로필에서 fill_down 켬
    _base_sheet(wb, "빈칸채움_끔", _fill_down_rows())    # 같은 원문, fill_down 끔

    _base_sheet(wb, "한칸경로", [["회원 > 로그인 > 아이디 찾기", "아이디 찾기 화면", "휴대폰 인증으로 찾기"]],
                header=["경로", "기능명", "설명"])

    _base_sheet(wb, "번호", [
        ["1", "회원", "회원 기능 묶음"],
        ["1.1", "로그인", "로그인 기능 묶음"],
        ["1.1.1", "아이디 찾기", "휴대폰 인증으로 찾기"],
        ["2", "주문", "주문 기능 묶음"],
        ["2.3.1", "주문 취소 사유", "사유 4종 중 선택"],  # 조상 2.3이 없음
    ], header=["번호", "이름", "설명"])

    _base_sheet(wb, "이어지는행", [
        ["V-001", "주문", "결제", "결제 화면", "결제 수단 선택"],
        [None, None, None, None, "쿠폰 적용 가능"],        # 중간 행: 기능명 칸 빈칸
        [None, None, None, "결제 확인 버튼", "최종 금액 확인"],
    ], header=["기능 ID", "분류1", "분류2", "기능명", "설명"])

    memo = _base_sheet(wb, "오른쪽메모", [
        ["주문", "장바구니", "담기", "담기 버튼", "옵션 선택 후 담기"],
        ["주문", "장바구니", "삭제", "삭제 버튼", "선택 항목 삭제"],
    ])
    _put(memo, {"H2": "메모: 다음 판에서 정리", "H3": "담당 확인 필요"})  # G열을 비우고 떨어진 덩어리

    _base_sheet(wb, "역할겹침", [["주문", "결제", "결제 화면", "결제 수단 선택", "카드·무통장 지원"]],
                header=["분류1", "분류2", "기능명", "설명", "상세 설명"])

    hidden = _base_sheet(wb, "숨김", [["회원", "로그인", None, "로그인 화면", "숨긴 시트의 항목"]])
    hidden.sheet_state = "hidden"

    two = _base_sheet(wb, "표두개", [
        ["주문", "결제", None, "결제 화면", "결제 수단 선택"],
        ["주문", "배송", None, "배송 조회 화면", "배송 단계 표시"],
    ])
    _rows(two, 7, 2, [BASE_HEADER, ["회원", "로그인", None, "로그인 화면", "아이디와 비밀번호 입력"]])  # 4~6행 빈 행 3줄

    indent = wb.create_sheet("들여쓰기")
    _rows(indent, 1, 2, [["기능 구조", "설명"], ["회원", "회원 기능"], ["로그인", "로그인 기능"], ["아이디 찾기", "휴대폰 인증"]])
    for r, level in ((2, 0), (3, 1), (4, 2)):
        indent.cell(r, 2).alignment = Alignment(indent=level)

    hrf = _base_sheet(wb, "숨김행수식", [
        ["회원", "로그인", None, "로그인 화면", "아이디와 비밀번호 입력"],
        ["회원", "로그인", None, "자동 로그인", "그냥 숨긴 행"],
        ["회원", "로그인", None, "로그인 유지 안내", "그룹 접기로 숨긴 행"],
        ["회원", "회원가입", None, "가입 화면", "약관 동의 후 가입"],
    ], header=BASE_HEADER + ["글자 수"])
    hrf.row_dimensions[3].hidden = True                      # 그냥 숨김
    hrf.row_dimensions[4].outline_level = 1                  # 그룹 접기
    hrf.row_dimensions[4].hidden = True
    hrf["G5"] = "=LEN(E5)"                                   # 계산값 없이 저장되는 수식

    skip = _base_sheet(wb, "숨김건너뜀", [
        ["회원", "로그인", None, "이전 로그인 화면", "숨긴 행 1"],
        ["회원", "로그인", None, "이전 가입 화면", "숨긴 행 2"],
        ["회원", "로그인", None, "로그인 화면", "아이디와 비밀번호 입력"],
        ["회원", "회원가입", None, "가입 화면", "약관 동의 후 가입"],
    ])
    skip.row_dimensions[2].hidden = True                     # 프로필에서 hidden_rows: skip
    skip.row_dimensions[3].hidden = True

    merged = _base_sheet(wb, "내용병합", [
        ["주문", "결제", "카드 결제", "카드 결제 버튼", "결제는 PG사 창에서 진행"],
        ["주문", "결제", "무통장 입금", "입금 계좌 안내", None],
        ["주문", "결제", "간편 결제", "간편 결제 버튼", None],
        ["주문", "배송", None, "배송 조회 화면", "배송 단계 표시"],
    ])
    merged.merge_cells("F2:F4")                              # 분류 열이 아닌 설명 열 3행 병합

    _base_sheet(wb, "분류패턴_2단", [
        ["주문", "장바구니", "장바구니 화면", "담은 목록 표시"],
        ["주문", "결제", "결제 화면", "결제 수단 선택"],
    ], header=["분류1", "분류2", "기능명", "설명"])
    _base_sheet(wb, "분류패턴_3단", [
        ["주문", "장바구니", "담기", "담기 버튼", "옵션 선택 후 담기"],
        ["주문", "결제", "카드 결제", "카드 결제 버튼", "국내 카드만 지원"],
    ])
    _base_sheet(wb, "분류열부족", [["주문", "결제", "결제 화면", "결제 수단 선택"]],
                header=["분류1", "분류2", "기능명", "설명"])  # 프로필은 분류1~3 선택자 3개

    note = _base_sheet(wb, "각주", [
        [1, "주문", "결제", None, "결제 화면", "결제 수단 선택"],
        [2, "주문", "배송", None, "배송 조회 화면", "배송 단계 표시"],
    ], header=["번호"] + BASE_HEADER)                         # 번호 열(B)은 역할 없음
    note["B5"] = "참고: 금액은 부가세 포함"                    # 빈 행 1줄 뒤, 번호 열에만 값, 끝 표시 글자 없음

    inflated = _base_sheet(wb, "부푼범위", [
        ["주문", "장바구니", None, "장바구니 화면", "담은 목록 표시", 3],
        ["주문", "결제", None, "결제 화면", "결제 수단 선택", 5],
        ["주문", "배송", None, "배송 조회 화면", "배송 단계 표시", 2],
    ], header=BASE_HEADER + ["개수"])
    inflated.cell(500, 2).font = Font(bold=True)            # 값 없이 서식만 → max_row가 부풂

    _base_sheet(wb, "표안빈열", [
        ["주문", "장바구니", "장바구니 화면", None, "담은 목록 표시"],
        ["주문", "결제", "결제 화면", "확인 필요", "결제 수단 선택"],
        ["주문", "배송", "배송 조회 화면", None, "배송 단계 표시"],
    ], header=["분류1", "분류2", "기능명", "메모", "설명"])

    diag = _base_sheet(wb, "대각선표", [
        ["주문", "결제", None, "결제 화면", "결제 수단 선택"],
        ["주문", "배송", None, "배송 조회 화면", "배송 단계 표시"],
    ])
    _rows(diag, 7, 8, [["구분", "값"], ["배송비", "3000"]])   # 4~6행 비우고 오른쪽 아래(H7)

    left = wb.create_sheet("왼쪽덩어리")
    _rows(left, 1, 4, [["분류1", "분류2", "기능명", "설명"],
                       ["주문", "결제", "결제 화면", "결제 수단 선택"],
                       ["주문", "배송", "배송 조회 화면", "배송 단계 표시"]])
    _put(left, {"B2": "임시 메모", "B3": "삭제 예정"})        # C열을 비우고 왼쪽 밖 덩어리

    chart = wb.create_chartsheet("차트")                    # 칸이 없는 차트 시트
    bar = BarChart()
    bar.add_data(Reference(inflated, min_col=7, min_row=1, max_row=4), titles_from_data=True)
    chart.add_chart(bar)

    wb.create_sheet("빈시트")                                # 아무 칸도 없음

    # 표여러개: 한 시트에 표 5개 — 제목 행(A열에만 값) + 다음 행 A열 「번호」 머리글일 때만 표 시작.
    # 표1 제목은 A1:D1 병합, 표2 끝 「※」(표마다 stop_markers), 표3은 분류 열이 없어 그 표만 건너뜀,
    # 표4·5는 머리글이 같고 빈 행 하나로 붙어 있음
    multi = wb.create_sheet("표여러개")
    head = ["번호", "분류1", "기능명", "설명"]
    _rows(multi, 1, 1, [
        ["주문 정책"], head, [1, "주문", "주문 취소", "결제 전에만 취소"], [2, None, "주문 변경", "배송 전에만 변경"], [],
        ["배송 정책"], head, [1, "배송", "배송 조회", "단계별 표시"], ["※ 도서 산간 지역 제외"], [],
        ["환불 정책"], ["번호", "기능명", "설명"], [1, "환불 신청", "수령 후 7일 이내"], [],
        ["가입 정책"], head, [1, "회원", "회원가입", "약관 동의 필수"], [],
        ["탈퇴 정책"], head, [1, "회원", "회원탈퇴", "탈퇴 사유 선택"],
    ])
    multi.merge_cells("A1:D1")
    wb.save(path)


# ---------- 미지원 형식 파일 ----------

OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")  # 옛 엑셀(.xls) 형식의 파일 앞부분


def make_ole_file(path):
    """openpyxl로는 옛 형식을 만들 수 없어 파일 앞부분 바이트를 직접 쓴다."""
    path.write_bytes(OLE_SIGNATURE + bytes(504))


MAKERS = {
    "가상몰_요구사항정의서.xlsx": make_req,
    "가상몰_기능정의서_v1.2.xlsx": make_fo,
    "가상몰_관리자기능정의서.xlsx": make_bo,
    "가상몰_정책정의서.xlsx": make_policy,
    "가상몰_테스트케이스.xlsx": make_test,
    "가상몰_변형.xlsx": make_variation,
    "가상몰_옛형식.xls": make_ole_file,        # .xls 미지원
    "가상몰_가짜확장자.xlsx": make_ole_file,   # U6: 확장자는 .xlsx인데 옛 형식
}


def generate(out_dir):
    """out_dir에 합성 엑셀을 모두 만들고 만든 파일 경로 목록을 돌려준다."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, make in MAKERS.items():
        path = out_dir / name
        make(path)
        paths.append(path)
    return paths


if __name__ == "__main__":
    generate(Path(__file__).parent / "generated")
