"""hierarchy ① columns(병합 복원·fill_down), none, 시트 이름 루트.
split_cell·numbering은 12, indent·parent_id는 이름만 → 미구현.
"""


import openpyxl
import pytest

from fixtures import expected as E
from si_prep import boundary, hierarchy
from si_prep.config import HIERARCHY_METHODS, load_project, load_profile, sheet_profile
from si_prep.issues import Issues
from si_prep.layout import build_table, build_tables, select_sheets
from si_prep.normalize import compare_key, path_text
from si_prep.reader import open_workbook
from si_prep.registry import Unimplemented
from conftest import SHOP, fill, make_profile



def _profile(**over):
    return make_profile({"hierarchy": [{"header_pattern": r"^분류\d$", "all": True}], "title": {"header_contains": "기능명"}},
                        **over)


def _wb(tmp_path, sheets):
    """{시트 이름: (행 목록, 병합 목록)} → 통합 문서"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (rows, merges) in sheets.items():
        fill(wb.create_sheet(name), rows, merges)
    path = tmp_path / "h.xlsx"
    wb.save(path)
    return open_workbook(path, Issues())


def _paths(wb, sheet, p, issues=None):
    issues = Issues() if issues is None else issues
    sp = sheet_profile(p, sheet)
    table = build_table(wb.sheet(sheet), sp, "h.xlsx", issues)
    groups = boundary.by_row(table, sp, issues)
    return hierarchy.REGISTRY.get(sp.hierarchy.method)(table, groups, sp, issues)


def _vals(steps):
    return [(s.value, s.basis, s.col) for s in steps]


# ---------- 등록표 ----------

def test_registry():
    assert sorted(HIERARCHY_METHODS) == ["columns", "indent", "none", "numbering", "parent_id", "split_cell"]
    assert hierarchy.REGISTRY.implemented() == ["columns", "none"]
    for name in ("split_cell", "numbering", "indent", "parent_id"):
        with pytest.raises(Unimplemented):
            hierarchy.REGISTRY.get(name)


# ---------- columns ----------

HEAD = ["분류1", "분류2", "분류3", "분류4", "기능명"]


def test_merge_restored_steps(tmp_path):
    wb = _wb(tmp_path, {"S": ([HEAD, ["가", "나", "다", None, "A"], [None, None, "라", None, "B"]], ["A2:A3", "B2:B3"])})
    paths = _paths(wb, "S", _profile())
    assert _vals(paths[1]) == [("가", "merge_restored", "A"), ("나", "merge_restored", "B"), ("라", "cell", "C")]


def test_trailing_empty_steps_dropped_without_warning(tmp_path):
    # 끝쪽 빈 칸(4단계만 빈 행 등)은 경로를 거기서 끝내고 경고 없음
    wb = _wb(tmp_path, {"S": ([HEAD, ["가", "나", "다", None, "A"], ["가", None, None, None, "B"]], [])})
    issues = Issues()
    paths = _paths(wb, "S", _profile(), issues)
    assert [len(p) for p in paths] == [3, 1] and len(issues) == 0


def test_middle_empty_step_is_empty_with_warning(tmp_path):
    wb = _wb(tmp_path, {"S": ([HEAD, ["가", None, "다", None, "A"]], [])})
    issues = Issues()
    [steps] = _paths(wb, "S", _profile(), issues)
    assert _vals(steps) == [("가", "cell", "A"), (None, "empty", "B"), ("다", "cell", "C")]
    assert path_text([s.value for s in steps]) == "가 › (빈칸) › 다"
    [w] = issues.items
    assert w.kind == "merge_unrestored" and [(l.rows.start, l.cols) for l in w.locations] == [(2, ["B"])]


FILL = [HEAD[:3] + ["기능명"],
        ["회원", "로그인", None, "A"],
        [None, None, "아이디 찾기", "B"],
        [None, "회원가입", None, "C"],
        [None, "회원탈퇴", None, "D"],
        [None, None, None, "E"]]


def test_fill_down_on_and_off_five_rows(tmp_path):
    wb = _wb(tmp_path, {"S": (FILL, [])})
    on = _paths(wb, "S", _profile(hierarchy={"method": "columns", "fill_down": True}))
    assert [_vals(p) for p in on] == [
        [("회원", "cell", "A"), ("로그인", "cell", "B")],
        [("회원", "filled_down", "A"), ("로그인", "filled_down", "B"), ("아이디 찾기", "cell", "C")],
        [("회원", "filled_down", "A"), ("회원가입", "cell", "B")],          # 새 값 아래 단계는 채우지 않음
        [("회원", "filled_down", "A"), ("회원탈퇴", "cell", "B")],
        [("회원", "filled_down", "A"), ("회원탈퇴", "filled_down", "B")],   # 분류 칸 전부 빈 행: 위 경로 그대로
    ]
    issues = Issues()
    off = _paths(wb, "S", _profile(), issues)
    assert [_vals(p) for p in off][1:] == [
        [(None, "empty", "A"), (None, "empty", "B"), ("아이디 찾기", "cell", "C")],
        [(None, "empty", "A"), ("회원가입", "cell", "B")],
        [(None, "empty", "A"), ("회원탈퇴", "cell", "B")],
        [(None, "empty", "A"), (None, "empty", "B"), (None, "empty", "C")],  # 단계 수 = 분류 열 수
    ]
    assert path_text([s.value for s in off[4]]) == "(빈칸) › (빈칸) › (빈칸)"
    assert [l.cols for w in issues.by_kind("merge_unrestored") for l in w.locations] == \
        [["A", "B"], ["A"], ["A"], ["A", "B", "C"]]


def test_fill_down_from_merged_value_above(tmp_path):
    # 위 행 값이 병합 복원이어도 그 값(=위 항목의 경로 값)으로 채운다
    rows = [HEAD[:3] + ["기능명"], ["가", "나", None, "A"], [None, "다", None, "B"], [None, None, "라", "C"]]
    wb = _wb(tmp_path, {"S": (rows, ["A2:A3"])})
    on = _paths(wb, "S", _profile(hierarchy={"method": "columns", "fill_down": True}))
    assert _vals(on[2]) == [("가", "filled_down", "A"), ("다", "filled_down", "B"), ("라", "cell", "C")]


# ---------- 시트 이름 루트·none ----------

def test_sheet_as_root_with_separators(tmp_path):
    wb = _wb(tmp_path, {"주문 - 결제": ([["정책명"], ["가"]], [])})
    p = _profile(columns={"title": {"header_contains": "정책명"}}, hierarchy={"method": "none"},
                 sheet_as_root=True, root_separators=[" - "])
    [steps] = _paths(wb, "주문 - 결제", p)
    assert _vals(steps) == [("주문", "sheet_name", None), ("결제", "sheet_name", None)]


def test_none_without_root_is_empty_path(tmp_path):
    wb = _wb(tmp_path, {"S": ([["이름"], ["가"]], [])})
    p = _profile(columns={"title": {"header_contains": "이름"}}, hierarchy={"method": "none"})
    assert _paths(wb, "S", p) == [[]]


def test_sheet_name_equal_to_first_step_matches_overview(tmp_path):
    # 메뉴 시트 이름과 1단계 값이 같을 때(시트 「가게」, 1단계 「가게」):
    # 메뉴 시트 경로 가게 › 가게 › 물건과 통합 시트 경로 가게 › 가게 › 물건·가게 › 물건이 같은 글자·같은 열쇠
    menu = ([["분류1", "분류2", "기능명"], ["가게", "물건", "A"]], [])
    overview = ([["분류1", "분류2", "분류3", "기능명"], ["가게", "가게", "물건", "A"], ["가게", "물건", None, "B"]], [])
    wb = _wb(tmp_path, {"가게": menu, "통합": overview})
    p = _profile(sheet_as_root=True, overview_sheets=["통합"],
                 sheet_overrides=[{"sheets": ["통합"], "sheet_as_root": False}])
    [m] = _paths(wb, "가게", p)
    o = _paths(wb, "통합", p)
    assert _vals(m) == [("가게", "sheet_name", None), ("가게", "cell", "A"), ("물건", "cell", "B")]
    texts = {path_text([s.value for s in x]) for x in [m] + o}
    keys = {compare_key([s.value for s in x]) for x in [m] + o}
    assert texts == {"가게 › 물건"} and keys == {("가게", "물건")}


# ---------- 합성 프로젝트: expected.py(FINAL)의 경로와 같음 ----------

def test_shop_paths_match_expected():
    config = load_project(SHOP)
    excluded = set(E.BEFORE_1008["excluded_sheets"])
    got, got_warn, processed = {}, set(), set()
    for doc in config.documents:
        profile = load_profile(config.profile_path(doc))
        for f in doc.files:
            wb = open_workbook(config.input_path / f, Issues())
            if wb is None:
                continue
            for grid in select_sheets(wb, profile, Issues()):
                sp = sheet_profile(profile, grid.name)
                if (f, grid.name) in excluded:
                    continue
                try:
                    split = boundary.REGISTRY.get(sp.item_boundary.method)
                    make = hierarchy.REGISTRY.get(sp.hierarchy.method)
                except Unimplemented:
                    continue
                issues = Issues()
                for table in build_tables(grid, sp, f, issues):   # 한 시트 여러 표면 표마다
                    processed.add((f, grid.name))
                    groups = split(table, sp, issues)
                    for g, steps in zip(groups, make(table, groups, sp, issues)):
                        got[(f, grid.name, g.start)] = [(s.value, s.basis, s.col) for s in steps]
                got_warn |= {(f, grid.name, l.rows.start, tuple(l.cols or ())) for w in issues.by_kind("merge_unrestored")
                             for l in w.locations}
    # 처리한 시트의 기대 항목은 빠짐없이 나와야 한다
    want = {(i["file"], i["sheet"], i["rows"][0]): [tuple(p) for p in i["path"]] for i in E.FINAL_ITEMS
            if (i["file"], i["sheet"]) in processed}
    assert got == want
    covered = {(f, s) for f, s, _ in got}
    for sheet in ("회원", "요구사항", "주문 - 배송", "테스트케이스", "빈칸채움", "빈칸채움_끔"):
        assert any(s == sheet for _, s in covered), sheet
    want_warn = {(w["file"], w["sheet"], a, tuple(c or ())) for w in E.FINAL_WARNINGS if w["kind"] == "merge_unrestored"
                 for _, a, _, c in w["locations"]}
    assert got_warn == want_warn


def test_fill_down_when_row_has_only_merge_restored_values(tmp_path):
    # 원문 새 값 없이 병합 복원 값만 있는 행도 아래 빈 단계를 위 항목에서 채운다
    rows = [HEAD[:3] + ["기능명"], ["가", "나", None, "A"], [None, None, None, "B"]]
    wb = _wb(tmp_path, {"S": (rows, ["A2:A3"])})
    on = _paths(wb, "S", _profile(hierarchy={"method": "columns", "fill_down": True}))
    assert _vals(on[1]) == [("가", "merge_restored", "A"), ("나", "filled_down", "B")]


def test_fill_down_does_not_fill_merge_with_empty_start(tmp_path):
    # 시작 칸이 빈 병합(A3:A4)은 「병합 아닌 빈칸」이 아니므로 채우지 않는다 → empty + merge_unrestored.
    # 같은 행의 병합 아닌 빈칸(B3)은 채운다
    rows = [HEAD[:3] + ["기능명"], ["가", "나", "다", "A"], [None, None, "마", "B"], [None, "바", "사", "C"]]
    wb = _wb(tmp_path, {"S": (rows, ["A3:A4"])})
    issues = Issues()
    on = _paths(wb, "S", _profile(hierarchy={"method": "columns", "fill_down": True}), issues)
    assert _vals(on[1]) == [(None, "empty", "A"), ("나", "filled_down", "B"), ("마", "cell", "C")]
    assert [l.cols for w in issues.by_kind("merge_unrestored") for l in w.locations][0] == ["A"]


# ---------- blank_means_above: 빈칸 = 위와 같음 사람 선언 ----------

BLANK_ROWS = [HEAD, ["가", "나", "다", None, "A"], [None, None, "라", None, "B"], [None, "마", None, None, "C"],
              [None, None, None, None, "D"]]


def _declared(**extra):
    return {"method": "columns", "fill_down": True, "blank_means_above": True, **extra}


def test_blank_means_above_marks_filled_steps_form_filled(tmp_path):
    wb = _wb(tmp_path, {"S": (BLANK_ROWS, [])})
    paths = _paths(wb, "S", _profile(hierarchy=_declared()))
    assert _vals(paths[1]) == [("가", "form_filled", "A"), ("나", "form_filled", "B"), ("라", "cell", "C")]
    assert _vals(paths[2]) == [("가", "form_filled", "A"), ("마", "cell", "B")]      # 새 값 아래 단계는 채우지 않음
    assert [s.basis for s in paths[3]] == ["filled_down", "filled_down"]           # 분류 칸 전부 빈 행은 추정 그대로


def test_without_declaration_still_filled_down(tmp_path):
    wb = _wb(tmp_path, {"S": (BLANK_ROWS, [])})
    paths = _paths(wb, "S", _profile(hierarchy={"method": "columns", "fill_down": True}))
    assert [s.basis for s in paths[1]] == ["filled_down", "filled_down", "cell"]


def test_blank_means_above_keeps_merge_restored(tmp_path):
    wb = _wb(tmp_path, {"S": ([HEAD, ["가", "나", "다", None, "A"], [None, None, "라", None, "B"]], ["A2:A3"])})
    paths = _paths(wb, "S", _profile(hierarchy=_declared()))
    assert _vals(paths[1]) == [("가", "merge_restored", "A"), ("나", "form_filled", "B"), ("라", "cell", "C")]


def test_blank_means_above_per_sheet_override(tmp_path):
    wb = _wb(tmp_path, {"S": (BLANK_ROWS, []), "T": (BLANK_ROWS, [])})
    p = _profile(hierarchy={"method": "columns", "fill_down": True},
                 sheet_overrides=[{"sheets": ["T"], "hierarchy": _declared()}])
    assert _paths(wb, "S", p)[1][0].basis == "filled_down"
    assert _paths(wb, "T", p)[1][0].basis == "form_filled"


def test_blank_means_above_needs_fill_down():
    from pydantic import ValidationError
    for hier in ({"method": "columns", "blank_means_above": True},
                 {"method": "columns", "fill_down": False, "blank_means_above": True}):
        with pytest.raises(ValidationError, match="fill_down"):
            _profile(hierarchy=hier)


# ---------- 12절 숫자 = by_columns 실제 채움 ----------

# 채울 수 없는 앞쪽 빈칸: 첫 항목(위 경로 없음), 시작 칸이 빈 병합(A3:A4) 아래의 빈칸, 위 단계가 빈 채로 이어지는 빈칸.
# 채울 수 있는 빈칸: 병합 복원 값(B6:B7) 아래의 빈칸, 분류 칸 전부 빈 행
UNFILLABLE = [HEAD[:3] + ["기능명"],
              [None, "나", "다", "A"],          # 첫 항목: A 빈칸인데 위 경로 없음 → 못 채움
              [None, None, None, "B"],          # 전부 빈 행(A3은 시작 칸이 빈 병합 A3:A4): B·C 채움
              [None, "라", "마", "C"],          # A4: 병합 칸이라 앞쪽 빈칸 아님
              [None, None, "바", "D"],          # A 빈칸: 위 A가 빈 병합이라 못 채움, B는 채움
              ["사", "아", None, "E"],          # B6:B7 병합
              [None, None, "자", "F"],          # A 채움(사), B는 병합 복원
              [None, None, "차", "G"],          # A·B 채움(병합 복원 값 아래)
              [None, None, None, "H"]]          # 전부 빈 행: 채움


def _table(wb, sheet, p):
    issues = Issues()
    sp = sheet_profile(p, sheet)
    table = build_table(wb.sheet(sheet), sp, "h.xlsx", issues)
    return table, boundary.by_row(table, sp, issues), sp


def _rows_with(paths, basis):
    return sum(any(s.basis == basis for s in p) for p in paths)


def test_blank_counts_match_filled_down_items(tmp_path):
    wb = _wb(tmp_path, {"S": (UNFILLABLE, ["A3:A4", "B6:B7"])})
    on = _profile(hierarchy={"method": "columns", "fill_down": True})
    paths = _paths(wb, "S", on)
    table, groups, _ = _table(wb, "S", on)
    b = hierarchy.blank_counts(table, groups)
    filled = _rows_with(paths, "filled_down")
    assert b.front_rows + b.all_blank_rows == filled == 5
    assert b.all_blank_rows == 2
    # 앞쪽 빈칸 = 실제로 채운 칸(전부 빈 행 제외)
    assert b.blank_front == sum(s.basis == "filled_down" for p, g in zip(paths, groups) for s in p
                                if any(table.grid.cell(g.rep, c).value is not None for c in table.roles["hierarchy"]))


def test_blank_counts_same_when_fill_down_off(tmp_path):
    # fill_down 끔 시트도 「켰다면 채울 칸」을 같은 판정으로 센다
    wb = _wb(tmp_path, {"S": (UNFILLABLE, ["A3:A4", "B6:B7"])})
    on = _table(wb, "S", _profile(hierarchy={"method": "columns", "fill_down": True}))
    off = _table(wb, "S", _profile())
    assert hierarchy.blank_counts(*off[:2]) == hierarchy.blank_counts(*on[:2])


def test_blank_counts_match_declared_sheet(tmp_path):
    wb = _wb(tmp_path, {"S": (UNFILLABLE, ["A3:A4", "B6:B7"])})
    p = _profile(hierarchy=_declared())
    paths = _paths(wb, "S", p)
    b = hierarchy.blank_counts(*_table(wb, "S", p)[:2])
    assert b.front_rows == _rows_with(paths, "form_filled")
    assert b.all_blank_rows == _rows_with(paths, "filled_down")


# ---------- B1 테두리 묶음: border_groups ----------

def _border_wb(tmp_path, rows, lines, merges=(), name="S"):
    """lines: {(행, 열 문자): "top"|"bottom"} 가로 선. 테두리만 있는 빈칸도 저장된다"""
    from openpyxl.styles import Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = name
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
    for (r, col), side in lines.items():
        ws[f"{col}{r}"].border = Border(**{side: Side(style="thin")})
    for m in merges:
        ws.merge_cells(m)
    path = tmp_path / "b.xlsx"
    wb.save(path)
    return open_workbook(path, Issues())


def _bordered(**extra):
    return _declared(border_groups=True, **extra)


# 분류1(A)은 테두리 묶음 양식: 2~3행 묶음(A2 값), 4행에서 새 묶음(위 선)인데 첫 칸 A4가 빔, 5행은 그 묶음 안
BORDER_ROWS = [HEAD[:2] + ["기능명"], ["가", "나", "A"], [None, "다", "B"], [None, "라", "C"], [None, "마", "D"]]
BORDER_LINES = {(2, "A"): "top", (4, "A"): "top", (3, "B"): "top", (4, "B"): "top", (5, "B"): "top"}


def test_border_groups_needs_blank_means_above():
    from pydantic import ValidationError
    for hier in ({"method": "columns", "fill_down": True, "border_groups": True},
                 {"method": "columns", "border_groups": True}):
        with pytest.raises(ValidationError, match="blank_means_above"):
            _profile(hierarchy=hier)
    assert _profile(hierarchy=_bordered()).hierarchy.border_groups is True
    assert _profile(hierarchy=_declared()).hierarchy.border_groups is False


def test_reader_line_above(tmp_path):
    # 위 칸의 아래 선이나 이 칸의 위 선이 있으면 「위에 선」. 값 없는 칸의 테두리도 읽는다
    wb = _border_wb(tmp_path, [["x"], [None], [None], [None]], {(2, "A"): "top", (3, "A"): "bottom"})
    g = wb.sheet("S")
    assert [g.line_above(r, "A") for r in (1, 2, 3, 4)] == [False, True, False, True]
    assert g.line_above(2, "B") is False


def test_border_groups_fill_inside_group_only(tmp_path):
    wb = _border_wb(tmp_path, BORDER_ROWS, BORDER_LINES)
    issues = Issues()
    paths = _paths(wb, "S", _profile(hierarchy=_bordered()), issues)
    assert _vals(paths[1]) == [("가", "form_filled", "A"), ("다", "cell", "B")]       # 같은 묶음 안 → 채움
    assert _vals(paths[2]) == [(None, "empty", "A"), ("라", "cell", "B")]             # 새 묶음 첫 칸이 빔 → 안 채움
    assert _vals(paths[3]) == [(None, "empty", "A"), ("마", "cell", "B")]             # 위가 비어 채울 값 없음
    [w, w2] = issues.by_kind("merge_unrestored")
    assert "테두리" in w.message and "새 묶음" in w.message and "A4" in w.message
    assert [(l.rows.start, l.cols) for l in w.locations] == [(4, ["A"])]
    assert "테두리" not in w2.message                                                 # 5행: 보통 빈 단계 경고


def test_without_border_groups_declared_fills_everything(tmp_path):
    # 같은 시트, border_groups 없이 선언만: 16c 그대로(새 묶음 첫 칸도 채움)
    wb = _border_wb(tmp_path, BORDER_ROWS, BORDER_LINES)
    paths = _paths(wb, "S", _profile(hierarchy=_declared()))
    assert [p[0].basis for p in paths[1:]] == ["form_filled"] * 3


def test_border_groups_column_without_lines_falls_back(tmp_path):
    # 분류 열에 가로 선이 하나도 없으면 테두리 묶음만 끄고 선언대로 채움 + 「테두리로 묶음을 구분할 수 없음」 경고 1건
    wb = _border_wb(tmp_path, BORDER_ROWS, {(3, "B"): "top"})
    issues = Issues()
    paths = _paths(wb, "S", _profile(hierarchy=_bordered()), issues)
    assert [p[0].basis for p in paths[1:]] == ["form_filled"] * 3
    ws = [w for w in issues if "테두리로 묶음을 구분할 수 없" in w.message]
    assert len(ws) == 1 and "A" in ws[0].message


def test_border_line_under_header_only_is_no_lines(tmp_path):
    # 머리글 바로 아래 선(첫 데이터 행의 위 선)만 있으면 묶음 구분 선이 없는 것
    wb = _border_wb(tmp_path, BORDER_ROWS, {(2, "A"): "top", (2, "B"): "top"})
    issues = Issues()
    _paths(wb, "S", _profile(hierarchy=_bordered()), issues)
    assert any("테두리로 묶음을 구분할 수 없" in w.message for w in issues)


def test_border_group_spans_multi_row_item(tmp_path):
    # 여러 행 항목(이어지는 행)이 끼어도 위 항목 대표 행과 이 행 사이 어디든 선이 있으면 새 묶음
    rows = [HEAD[:2] + ["기능명"], ["가", "나", "A"], [None, "다", "B"], [None, None, None], [None, "라", "C"]]
    wb = _border_wb(tmp_path, rows, {(2, "A"): "top", (4, "A"): "top", (3, "B"): "top", (5, "B"): "top"})
    paths = _paths(wb, "S", _profile(hierarchy=_bordered()))
    assert [p[0].basis for p in paths[1:]] == ["form_filled", "empty"]


def test_blank_counts_border_inner(tmp_path):
    # 12절 「테두리 묶음 안 빈칸」: 앞쪽 빈칸 중 위 항목과 사이에 가로 선이 없는 칸(선언과 상관없이 셈)
    wb = _border_wb(tmp_path, BORDER_ROWS, BORDER_LINES)
    for p in (_profile(hierarchy={"method": "columns", "fill_down": True}), _profile(hierarchy=_bordered())):
        b = hierarchy.blank_counts(*_table(wb, "S", p)[:2])
        assert (b.blank_front, b.border_inner) == (3, 2)        # 3·4·5행 채움, 4행만 위에 선


def test_border_cut_trailing_step_kept_and_warned(tmp_path):
    # 리뷰 지적: 분류 칸이 전부 빈 행에서 맨 뒤 단계 위에만 선이 있으면, 그 단계는 버리지 않고
    # empty로 남기고 「테두리로 새 묶음」 경고(경로만 짧아지고 경고가 없으면 안 됨)
    rows = [HEAD[:3] + ["기능명"], ["가", "나", "다", "A"], [None, None, None, "B"]]
    lines = {(2, c): "top" for c in "ABC"} | {(3, "C"): "top", (4, "A"): "top", (4, "B"): "top"}
    wb = _border_wb(tmp_path, rows + [["라", "마", "바", "C"]], lines)
    issues = Issues()
    paths = _paths(wb, "S", _profile(hierarchy=_bordered()), issues)
    assert _vals(paths[1]) == [("가", "filled_down", "A"), ("나", "filled_down", "B"), (None, "empty", "C")]
    [w] = issues.by_kind("merge_unrestored")
    assert "C3" in w.message and "새 묶음" in w.message


def test_border_line_under_merged_block_read(tmp_path):
    # 병합 칸(A2:A3) 아래 칸의 아래 선이 병합 뒤에도 읽힌다 → 그 아래 빈칸은 새 묶음
    from openpyxl.styles import Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r, row in enumerate([HEAD[:2] + ["기능명"], ["가", "나", "A"], [None, "다", "B"], [None, "라", "C"]], start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
    ws["A3"].border = Border(bottom=Side(style="thin"))
    ws.merge_cells("A2:A3")
    path = tmp_path / "m.xlsx"
    wb.save(path)
    g = open_workbook(path, Issues()).sheet("S")
    assert g.line_above(4, "A") is True
