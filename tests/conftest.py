"""테스트 시작 시 합성 엑셀을 tests/fixtures/generated/에 만든다. 여러 테스트 파일이 같이 쓰는 경로·도우미도 여기 둔다."""

from pathlib import Path

import openpyxl
import pytest

from fixtures import make_fixtures
from si_prep.config import Profile

GENERATED_DIR = Path(__file__).parent / "fixtures" / "generated"
SHOP_DIR = Path(__file__).parent / "fixtures" / "configs" / "shop"
SHOP = SHOP_DIR / "project.yaml"


def make_profile(columns, /, **over):
    """한 시트 시험용 프로필: 1행 머리글, 2행부터 데이터, 분류는 columns 방식, 한 행 = 한 항목. over로 덮어씀(columns= 포함)"""
    data = {"id": "x", "header_rows": [1], "data_start_row": 2, "columns": columns,
            "hierarchy": {"method": "columns"}, "item_boundary": {"method": "row"}}
    data.update(over)
    return Profile.model_validate(data)


def fill(ws, rows, merges=()):
    """시트에 행 목록(None은 빈 칸)을 1행·A열부터 쓰고 병합한다"""
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
    for m in merges:
        ws.merge_cells(m)


def write_book(path, rows, merges=(), sheet="S", hidden_rows=(), hidden_cols=(), extra=None):
    """시트 하나짜리 합성 엑셀을 path에 저장하고 path를 돌려준다. extra(wb)로 시트를 더 붙일 수 있다"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    fill(ws, rows, merges)
    for r in hidden_rows:
        ws.row_dimensions[r].hidden = True
    for col in hidden_cols:
        ws.column_dimensions[col].hidden = True
    if extra:
        extra(wb)
    wb.save(path)
    return path


@pytest.fixture(scope="session", autouse=True)
def generated_dir():
    make_fixtures.generate(GENERATED_DIR)
    return GENERATED_DIR
